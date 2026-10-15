"""Drives the real tkinter window (needs a display: run under xvfb-run on Linux)."""
import os
import sys
import time

import openpyxl
import pytest

tk = pytest.importorskip("tkinter")
if sys.platform != "win32" and not os.environ.get("DISPLAY"):
    pytest.skip("no display", allow_module_level=True)

import numpy as np

from quiniela_scanner import gui
from quiniela_scanner.camera import ImageSource
from quiniela_scanner.config import Config
from quiniela_scanner.excel_sink import OpenpyxlBackend
from quiniela_scanner.name_reader import NameResult
from helpers import draw_slip, make_pool_workbook, on_webcam_canvas
import cv2

PICKS = list("LVLLELLVE")


class FakeReader:
    def __init__(self, res): self.res = res
    def read(self, name_crop, title_crop=None):
        time.sleep(0.2)
        return self.res


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(gui, "app_dir", lambda: tmp_path)
    monkeypatch.setattr(gui.messagebox, "askyesno", lambda *a, **k: True)
    monkeypatch.setattr(gui.messagebox, "showerror", lambda *a, **k: (_ for _ in ()).throw(AssertionError(a)))
    photo = tmp_path / "slip.jpg"
    cv2.imwrite(str(photo), on_webcam_canvas(draw_slip(PICKS, rng=np.random.default_rng(1)), scale=0.5, angle=4))
    return tmp_path, photo


def make_app(tmp_path, photo, wbname="JORNADA 12 PC.xlsx", reader=None, **cfgkw):
    xlsx = make_pool_workbook(tmp_path / wbname)
    cfg = Config()
    cfg.read_names_with_claude = False
    cfg.stable_frames = 4
    for k, v in cfgkw.items():
        setattr(cfg, k, v)
    src = ImageSource(str(photo), demo=True, hold=200, gap=60)
    app = gui.App(cfg, source=src, backend=OpenpyxlBackend(str(xlsx)))
    if reader is not None:
        app.reader = reader
    return app, xlsx


def pump(app, cond, timeout=15.0):
    end = time.time() + timeout
    while time.time() < end:
        app.update()
        if cond():
            return True
        time.sleep(0.02)
    return False


def test_manual_flow_type_name_and_press_enter(env):
    tmp, photo = env
    app, xlsx = make_app(tmp, photo)
    try:
        assert pump(app, lambda: app.pending is not None), "slip was never captured"
        assert [v.get() for v in app.pick_vars] == PICKS
        assert app.block_var.get().startswith("ENTRY PG1  N:") and app.ticket_var.get() == "251"
        assert app.tag_on.get() is False and app.tag_state.get() == "tag off"   # tagging starts OFF
        app.name_var.set("clemente o.")
        app.commit()                                                 # = pressing Enter
        ws = openpyxl.load_workbook(xlsx)["ENTRY PG1"]
        assert ws["N15"].value == "CLEMENTE O."                      # no tag: written as scanned
        assert [ws.cell(15, 15 + k).value for k in range(9)] == PICKS
        assert app.pending is None and app.ticket_var.get() == "252"
        assert len(app.tree.get_children()) == 1
        assert (tmp / "scans" / "log.csv").exists()
        app.undo()
        assert openpyxl.load_workbook(xlsx)["ENTRY PG1"]["N15"].value is None
        assert app.ticket_var.get() == "251"
    finally:
        app.destroy()


def test_auto_enter_when_claude_read_is_good_and_jornada_matches(env):
    tmp, photo = env
    app, xlsx = make_app(tmp, photo, wbname="JORNADA 12 PC.xlsx", auto_enter_delay_s=0.5,
                         reader=FakeReader(NameResult("CLEMENTE O.", "12", True)))
    try:
        app.tag_var.set("RRR"); app.tag_on.set(True)                  # group tag switched on before the slip arrives
        assert pump(app, lambda: len(app.tree.get_children()) == 1, 20), "auto-enter never happened"
        ws = openpyxl.load_workbook(xlsx)["ENTRY PG1"]
        assert ws["N15"].value == "CLEMENTE O. RRR"
        assert [ws.cell(15, 15 + k).value for k in range(9)] == PICKS
    finally:
        app.destroy()


def test_jornada_mismatch_blocks_auto_enter(env):
    tmp, photo = env
    app, xlsx = make_app(tmp, photo, wbname="JORNADA 10 PC.xlsx", auto_enter_delay_s=0.3,
                         reader=FakeReader(NameResult("CLEMENTE O.", "12", True)))
    try:
        assert pump(app, lambda: app.pending is not None and app.pending.name_done)
        pump(app, lambda: False, 1.5)                                # give auto-enter every chance
        assert app.pending is not None and not app.tree.get_children()
        assert "JORNADA 12" in app.jor_note.get() and "JORNADA 10" in app.jor_note.get()
    finally:
        app.destroy()


def test_unsure_name_or_edit_blocks_auto_enter(env):
    tmp, photo = env
    app, _ = make_app(tmp, photo, auto_enter_delay_s=0.3, reader=FakeReader(NameResult("MAYBE", "12", False)))
    try:
        assert pump(app, lambda: app.pending is not None and app.pending.name_done)
        pump(app, lambda: False, 1.2)
        assert app.pending is not None and not app.tree.get_children()
    finally:
        app.destroy()


def test_flagged_rows_need_a_human(env):
    tmp, _ = env
    picks = ["L", "", "LV"] + list("ELLVE")[:6]
    photo = tmp / "flagged.jpg"
    cv2.imwrite(str(photo), on_webcam_canvas(draw_slip(picks, rng=np.random.default_rng(1)), scale=0.5, angle=4))
    app, xlsx = make_app(tmp, photo, auto_enter_delay_s=0.3, reader=FakeReader(NameResult("ANA", "12", True)))
    try:
        assert pump(app, lambda: app.pending is not None and app.pending.name_done)
        pump(app, lambda: False, 1.2)
        assert app.pending is not None and not app.tree.get_children()   # still waiting for the operator
        assert app.pick_vars[1].get() == "" and app.pick_vars[2].get() == ""
        app.pick_vars[1].set("E"); app.pick_edited(1)
        app.pick_vars[2].set("V"); app.pick_edited(2)
        app.commit()
        ws = openpyxl.load_workbook(xlsx)["ENTRY PG1"]
        assert [ws.cell(15, 15 + k).value for k in range(3)] == ["L", "E", "V"]
    finally:
        app.destroy()


def test_typed_ticket_number_is_honoured_and_manual_rows_are_respected(env):
    tmp, photo = env
    app, xlsx = make_app(tmp, photo)
    try:
        assert pump(app, lambda: app.pending is not None)
        # someone types a name into the next free row by hand while the slip is on screen
        wb = openpyxl.load_workbook(xlsx)
        wb["ENTRY PG1"]["N15"] = "TYPED BY HAND"
        wb.save(xlsx)
        app.name_var.set("ANA")
        app.commit()
        ws = openpyxl.load_workbook(xlsx)["ENTRY PG1"]
        assert ws["N15"].value == "TYPED BY HAND"                     # untouched
        assert ws["N16"].value == "ANA"                               # went to the next free row instead
        # now a deliberate ticket number
        assert pump(app, lambda: app.pending is not None, 25)
        app.ticket_var.set("260"); app.ticket_custom = True
        app.name_var.set("BETO")
        app.commit()
        ws = openpyxl.load_workbook(xlsx)["ENTRY PG1"]
        assert ws["N24"].value == "BETO"                              # ticket 260 -> row 5 + 19
    finally:
        app.destroy()


def test_group_tag_on_off_before_after_and_single_letter(env):
    tmp, photo = env
    app, xlsx = make_app(tmp, photo)
    names = lambda: openpyxl.load_workbook(xlsx)["ENTRY PG1"]
    try:
        assert pump(app, lambda: app.pending is not None)
        app.name_var.set("ana")
        assert app.final_var.get().endswith("ANA")                    # preview with the tag off
        app.tag_var.set("rrr"); app.tag_on.set(True)                  # switching on mid-slip...
        assert app.final_var.get() == "Will be written as:  ANA RRR"
        assert app.pending.touched                                    # ...never auto-writes, needs Enter
        assert "TAG ON" in app.tag_state.get()
        app.commit()
        assert names()["N15"].value == "ANA RRR"

        assert pump(app, lambda: app.pending is not None, 25)         # next slip: same tag, but BEFORE the name
        app.tag_pos.set("before")
        app.name_var.set("beto")
        app.commit()
        assert names()["N16"].value == "RRR BETO"

        assert pump(app, lambda: app.pending is not None, 25)         # next person has no group: switch off
        app.tag_on.set(False)
        app.name_var.set("carlos")
        app.commit()
        assert names()["N17"].value == "CARLOS"

        assert pump(app, lambda: app.pending is not None, 25)         # a one-letter tag, on a name that ends in E
        app.tag_var.set("e"); app.tag_pos.set("after"); app.tag_on.set(True)
        app.name_var.set("jose")
        app.commit()
        assert names()["N18"].value == "JOSE E"
    finally:
        app.destroy()


def test_tag_switch_starts_off_and_remembers_text_and_position(env):
    tmp, photo = env
    app, _ = make_app(tmp, photo, name_tag="RRR", name_tag_position="before")
    try:
        assert app.tag_var.get() == "RRR" and app.tag_pos.get() == "before"
        assert app.tag_on.get() is False
    finally:
        app.destroy()
