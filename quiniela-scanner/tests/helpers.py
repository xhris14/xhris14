"""Test fixtures: a pool workbook laid out like the real one, and synthetic slip photos."""
from __future__ import annotations

import cv2
import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter

TEAMS = [("NECAXA", "ATLAS"), ("ATLANTE", "PUMAS"), ("TIJUANA", "PUEBLA"), ("CHIVAS", "TIGRES"),
         ("SANTOS", "QUERETARO"), ("LEON", "AMERICA"), ("TOLUCA", "SAN LUIS"),
         ("CRUZ AZUL", "JUAREZ"), ("MONTERREY", "PACHUCA")]


def make_pool_workbook(path, filled=(10, 3, 0), first_tickets=(241, 481, 721), rows=52, suffixes=("RRR", "CH", ""),
                       hash_headers=(True, True, True)):
    """ENTRY PG1 with three side-by-side blocks:  # | NOMBRE | 1..9 | T   (header on row 4).
    hash_headers[b] False leaves the '#' cell of block b empty (the ticket numbers are still there, like the first block
    of the real sheet)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ENTRY PG1"
    wb.create_sheet("ENTRY PG2")
    wb.create_sheet("RESULTS")
    wb.create_sheet("SLIPS")
    ws["N3"] = "SOLO DIVERSION"
    start_col = 13                              # column M, like the real sheet
    for b, (first, nfilled, suf) in enumerate(zip(first_tickets, filled, suffixes)):
        c0 = start_col + b * 12                 # '#', NOMBRE, 9 picks, T -> 12 columns
        if hash_headers[b]:
            ws.cell(4, c0, "#")
        ws.cell(4, c0 + 1, "NOMBRE")
        for k in range(9):
            ws.cell(4, c0 + 2 + k, k + 1)
        ws.cell(4, c0 + 11, "T")
        for i in range(rows):
            r = 5 + i
            ws.cell(r, c0, first + i)
            pick_first = get_column_letter(c0 + 2)
            pick_last = get_column_letter(c0 + 10)
            ws.cell(r, c0 + 11, f'=COUNTIF({pick_first}{r}:{pick_last}{r},"V")')
            if i < nfilled:
                ws.cell(r, c0 + 1, f"PLAYER {i + 1} {suf}".strip())
                for k in range(9):
                    ws.cell(r, c0 + 2 + k, "LEV"[(i + k) % 3])
    wb.save(path)
    return path


# ----------------------------------------------------------------------------- slips
def draw_slip(picks, name="CLEMENTE O.", ink=(170, 60, 25), width=1000, rng=None, jitter=True):
    """Draw an upright slip like the real form. picks: list of 9 items, each '' / 'L' / 'E' / 'V' / 'LV' (two marks)."""
    rng = rng or np.random.default_rng(0)
    h = 1050
    img = np.full((h, width, 3), 238, np.uint8)
    blk = (25, 25, 25)
    x = [8, 72, 136, 200, 585, 990]            # vertical rules: border, L|E, E|V, V|team, home|away, border
    y0, ytitle, yhead = 8, 85, 160
    row_h = 80
    ys = [yhead + i * row_h for i in range(10)]    # 10 lines bound the 9 match rows
    ybot = ys[-1] + 90                              # NOMBRE row bottom
    t = 4
    cv2.rectangle(img, (x[0], y0), (x[-1], ybot), blk, t)
    cv2.line(img, (x[0], ytitle), (x[-1], ytitle), blk, t)
    for yy in ys:
        cv2.line(img, (x[0], yy), (x[-1], yy), blk, t)
    for xx in x[1:5]:
        cv2.line(img, (xx, ytitle), (xx, ys[-1]), blk, t)
    cv2.putText(img, "LIGA MX JORNADA 12", (480, 62), cv2.FONT_HERSHEY_DUPLEX, 1.1, blk, 3)
    for lab, xx in zip("LEV", (22, 88, 152)):
        cv2.putText(img, lab, (xx, 125), cv2.FONT_HERSHEY_DUPLEX, 0.9, blk, 3)
    for i, (home, away) in enumerate(TEAMS):
        base = ys[i] + 52
        cv2.putText(img, home, (215, base), cv2.FONT_HERSHEY_DUPLEX, 1.1, blk, 3)
        cv2.putText(img, away, (600, base + 10), cv2.FONT_HERSHEY_DUPLEX, 1.1, blk, 3)
    cv2.putText(img, "NOMBRE", (14, ybot - 28), cv2.FONT_HERSHEY_DUPLEX, 0.95, blk, 3)
    cv2.line(img, (170, ybot - 22), (900, ybot - 22), blk, 2)
    cv2.putText(img, name, (260, ybot - 30), cv2.FONT_HERSHEY_SCRIPT_COMPLEX, 2.2, ink, 4)
    cv2.putText(img, "ULTIMO DIA MIERCOLES EL 14", (190, ybot + 40), cv2.FONT_HERSHEY_DUPLEX, 0.9, blk, 3)
    for i, p in enumerate(picks):
        for letter in p:
            c = "LEV".index(letter)
            cx = (x[c] + x[c + 1]) // 2
            cy = (ys[i] + ys[i + 1]) // 2
            jx, jy = (rng.integers(-8, 9), rng.integers(-8, 9)) if jitter else (0, 0)
            s = 24
            cv2.line(img, (cx - s + jx, cy - s + jy), (cx + s + jx, cy + s + jy), ink, 4)
            cv2.line(img, (cx + s + jx, cy - s + jy), (cx - s + jx, cy + s + jy), ink, 4)
    return img


def on_webcam_canvas(slip, scale=0.5, angle=0.0, persp=0.0, bg=235, blur=0.0, noise=0.0, seed=0, flip=False):
    """Place the slip on a 1920x1080 'webcam' frame with rotation / perspective / blur / noise."""
    rng = np.random.default_rng(seed)
    if flip:
        slip = cv2.rotate(slip, cv2.ROTATE_180)
    s = cv2.resize(slip, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA)
    pad = int(0.3 * max(s.shape[:2]))                       # room to rotate without clipping the slip
    s = cv2.copyMakeBorder(s, pad, pad, pad, pad, cv2.BORDER_CONSTANT, value=(bg,) * 3)
    h, w = s.shape[:2]
    m = cv2.getRotationMatrix2D((w / 2, h / 2), angle, 1.0)
    s = cv2.warpAffine(s, m, (w, h), borderValue=(bg,) * 3)
    if persp:
        src = np.float32([[0, 0], [w, 0], [w, h], [0, h]])
        dst = src + rng.uniform(-persp * min(w, h), persp * min(w, h), src.shape).astype(np.float32)
        s = cv2.warpPerspective(s, cv2.getPerspectiveTransform(src, dst), (w, h), borderValue=(bg,) * 3)
    canvas = np.full((1080, 1920, 3), bg, np.uint8)
    ox, oy = (1920 - w) // 2, (1080 - h) // 2
    canvas[oy:oy + h, ox:ox + w] = s
    if blur:
        canvas = cv2.GaussianBlur(canvas, (0, 0), blur)
    if noise:
        canvas = np.clip(canvas + rng.normal(0, noise, canvas.shape), 0, 255).astype(np.uint8)
    return canvas


# ----------------------------------------------------------------------------- stacks of slips
_NOISE_BANKS: dict = {}


def _noise_bank(sigma, shape, n=4):
    """A few pre-drawn sensor-noise fields; each frame takes a random crop of a random one."""
    key = (sigma, shape)
    if key not in _NOISE_BANKS:
        rng = np.random.default_rng(1234)
        _NOISE_BANKS[key] = [np.rint(rng.normal(0, sigma, (shape[0] + 48, shape[1] + 48, 3))).astype(np.int16)
                             for _ in range(n)]
    return _NOISE_BANKS[key]


def draw_hand(canvas, x, y, angle=0.0, size=1.0, tone=(118, 152, 208)):
    """A skin-coloured hand (palm, four fingers, thumb, forearm) centred on (x, y); angle in degrees."""
    a = np.deg2rad(angle)
    rot = np.array([[np.cos(a), -np.sin(a)], [np.sin(a), np.cos(a)]])

    def pt(px, py):
        return tuple(int(v) for v in (rot @ np.array([px, py]) * size + (x, y)))
    t = tuple(int(c) for c in tone)
    shade = tuple(int(c * 0.9) for c in tone)
    cv2.fillPoly(canvas, [np.array([pt(-90, 40), pt(90, 40), pt(120, 700), pt(-120, 700)], np.int32)], shade, cv2.LINE_AA)
    cv2.ellipse(canvas, pt(0, 0), (int(95 * size), int(75 * size)), angle, 0, 360, t, -1, cv2.LINE_AA)
    for fx, ln in ((-66, 120), (-22, 140), (22, 135), (64, 105)):
        cv2.line(canvas, pt(fx, -50), pt(fx, -50 - ln), t, int(34 * size), cv2.LINE_AA)
    cv2.line(canvas, pt(-80, 0), pt(-150, -60), t, int(36 * size), cv2.LINE_AA)


def _paste(canvas, small, cx, cy, rot=0.0, sc=1.0, alpha=None):
    """Paste ``small`` (a slip image) centred on (cx, cy), rotated ``rot`` degrees and scaled ``sc``.
    With ``alpha`` (uint8, canvas sized) the pasted pixels are also marked there."""
    h, w = small.shape[:2]
    m = cv2.getRotationMatrix2D((w / 2, h / 2), rot, sc)
    m[0, 2] += cx - w / 2
    m[1, 2] += cy - h / 2
    pts = cv2.transform(np.float32([[[0, 0], [w, 0], [w, h], [0, h]]]), m)[0]
    x0, y0 = np.floor(pts.min(axis=0)).astype(int) - 2
    x1, y1 = np.ceil(pts.max(axis=0)).astype(int) + 2
    x0, y0, x1, y1 = max(x0, 0), max(y0, 0), min(x1, canvas.shape[1]), min(y1, canvas.shape[0])
    if x1 <= x0 or y1 <= y0:
        return
    m[0, 2] -= x0
    m[1, 2] -= y0
    size = (x1 - x0, y1 - y0)
    warped = cv2.warpAffine(small, m, size, flags=cv2.INTER_LINEAR)
    mask = cv2.warpAffine(np.full((h, w), 255, np.uint8), m, size, flags=cv2.INTER_NEAREST) > 0
    roi = canvas[y0:y1, x0:x1]
    roi[mask] = warped[mask]
    if alpha is not None:
        alpha[y0:y1, x0:x1][mask] = 255


def _over(static, layer, alpha, blur):
    """Composite a moving layer (slip + hand) over the static frame, motion-blurring only the layer."""
    x, y, w, h = cv2.boundingRect(alpha)
    if w == 0:
        return static
    pad = blur + 2
    x0, y0, x1, y1 = max(x - pad, 0), max(y - pad, 0), min(x + w + pad, 1920), min(y + h + pad, 1080)
    a = alpha[y0:y1, x0:x1].astype(np.float32) / 255
    col = layer[y0:y1, x0:x1].astype(np.float32) * a[..., None]
    if blur > 1:
        k = (int(blur), max(1, int(blur) // 2))
        a, col = cv2.blur(a, k), cv2.blur(col, k)
    out = static[y0:y1, x0:x1].astype(np.float32) * (1 - a[..., None]) + col
    static[y0:y1, x0:x1] = np.clip(out, 0, 255).astype(np.uint8)
    return static


class Stage:
    """A fixed 1080p camera over a stack of slips, as 30 fps frame generators.

    ``add(slip_image)`` registers a slip and returns its id; the sequence methods yield
    ``(frame, info)`` where info = dict(phase=..., top=id, i=frame number within the phase).
    Every frame gets independent sensor noise (``noise`` sigma), a global brightness flicker
    (+-``flicker``) and, for a slip lying still, an autofocus-breathing scale jitter (+-``breath``).
    """

    CX, CY = 960, 540
    WIN = (780, 1260)                         # noisy window (h, w) around the slip, to keep frames cheap

    def __init__(self, seed=0, scale=0.5, angle=4.0, noise=3.0, flicker=0.03, breath=0.015, bg=235):
        self.rng = np.random.default_rng(seed)
        self.scale, self.angle, self.noise, self.flicker, self.breath, self.bg = scale, angle, noise, flicker, breath, bg
        self.slips: list = []
        self.offsets: list = []

    def add(self, slip) -> int:
        self.slips.append(cv2.resize(slip, None, fx=self.scale, fy=self.scale, interpolation=cv2.INTER_AREA))
        # a neat stack is never perfectly aligned: each sheet sits a few px / a fraction of a degree off
        self.offsets.append((self.rng.uniform(-4, 4), self.rng.uniform(-4, 4), self.rng.uniform(-0.4, 0.4)))
        return len(self.slips) - 1

    # -- one frame -------------------------------------------------------------------
    def _compose(self, top=None, under=None, lift=0.0, wobble=0.0, hand=None, blur=0, gain=None):
        canvas = np.full((1080, 1920, 3), self.bg, np.uint8)
        moving = lift > 0 or hand is not None
        layer = np.zeros_like(canvas) if moving else None
        alpha = np.zeros(canvas.shape[:2], np.uint8) if moving else None
        if under is not None:
            ox, oy, orot = self.offsets[under]
            _paste(canvas, self.slips[under], self.CX + ox, self.CY + oy, self.angle + orot)
        cx = cy = 0.0
        if top is not None:
            ox, oy, orot = self.offsets[top]
            sc = 1.0 + (self.rng.uniform(-self.breath, self.breath) if lift == 0 else 0.0) + 0.10 * lift
            cx, cy = self.CX + ox + 1450 * lift, self.CY + oy - 800 * lift
            rot = self.angle + orot + 25.0 * lift
            if wobble:
                cx += self.rng.uniform(-wobble, wobble)
                cy += self.rng.uniform(-wobble, wobble)
                rot += self.rng.uniform(-wobble, wobble) / 8
            if lift > 0:
                _paste(layer, self.slips[top], cx, cy, rot, sc, alpha)
            else:
                _paste(canvas, self.slips[top], cx, cy, rot, sc)
        if hand is True and top is not None:
            draw_hand(layer, cx + 150, cy + 140, 200 - 25 * lift, 1.0)
        elif isinstance(hand, tuple):
            draw_hand(layer, hand[0], hand[1], hand[2] if len(hand) > 2 else 0.0, 1.0)
        if moving:
            if hand is not None:
                alpha[(layer != 0).any(axis=2)] = 255          # the hand (the slip was marked on paste)
            canvas = _over(canvas, layer, alpha, blur)
        h, w = self.WIN
        y0, x0 = (1080 - h) // 2, (1920 - w) // 2
        g = 1.0 + self.rng.uniform(-self.flicker, self.flicker) if gain is None else gain
        if g != 1.0:
            canvas = cv2.convertScaleAbs(canvas, alpha=g)
        if self.noise:
            bank = _noise_bank(self.noise, self.WIN)[int(self.rng.integers(4))]
            oy, ox = int(self.rng.integers(0, 48)), int(self.rng.integers(0, 48))
            win = canvas[y0:y0 + h, x0:x0 + w]
            canvas[y0:y0 + h, x0:x0 + w] = np.clip(win.astype(np.int16) + bank[oy:oy + h, ox:ox + w], 0, 255).astype(np.uint8)
        return canvas

    # -- sequences -------------------------------------------------------------------
    def hold(self, top, n=21, under=None, wobble=0.0, phase="hold"):
        """``top`` lies still (sensor noise, flicker, focus breathing) for n frames."""
        for i in range(n):
            yield self._compose(top, under, wobble=wobble), dict(phase=phase, top=top, i=i)

    def empty(self, n=5, under=None):
        for i in range(n):
            yield self._compose(None, under), dict(phase="empty", top=under, i=i)

    def lift(self, top, under=None, n=6, hand=True):
        """The top slip is picked up (a hand holds it) and carried out of the picture in n frames."""
        for i in range(n):
            u = ((i + 1) / n) ** 1.6
            blur = int(10 + 60 * u) if hand else 0
            yield self._compose(top, under, lift=u * 1.25, hand=hand or None, blur=blur), dict(phase="lift", top=top, i=i)

    def drop(self, top, under=None, n=5, hand=True):
        """Reverse of lift: the slip is carried in and put down on the stack."""
        for i in range(n):
            u = (1 - (i + 1) / n) ** 1.6
            blur = int(10 + 60 * u) if hand else 0
            yield self._compose(top, under, lift=u * 1.25, hand=(hand or None) if u > 0 else None, blur=blur), \
                dict(phase="drop", top=top, i=i)

    def hand_pass(self, n=8, top=None, under=None):
        """A hand sweeps across the camera view; the slip below (if any) does not move."""
        for i in range(n):
            x = -250 + (2400 * (i + 1)) / n
            yield self._compose(top, under, hand=(x, self.CY + 40, 90), blur=60), dict(phase="hand", top=top, i=i)

    def slide(self, top, n=12, step=20.0):
        """The slip keeps sliding sideways ``step`` px per frame: it never sits still."""
        for i in range(n):
            ox, oy, orot = self.offsets[top]
            canvas = np.full((1080, 1920, 3), self.bg, np.uint8)
            _paste(canvas, self.slips[top], self.CX + ox - n * step / 2 + i * step, self.CY + oy, self.angle + orot)
            yield canvas, dict(phase="slide", top=top, i=i)


def draw_slip_named(picks, name, seed=0, ink=(170, 60, 25), scale=2.2, thick=4, dx=0, dy=0):
    """draw_slip with the name drawn separately, so the 'handwriting' (size, weight, offset) can differ
    between two slips with the same picks; ``seed`` drives the mark jitter. name='' leaves the strip blank."""
    img = draw_slip(picks, name="", ink=ink, rng=np.random.default_rng(seed))
    if name:
        cv2.putText(img, name, (260 + dx, 940 + dy), cv2.FONT_HERSHEY_SCRIPT_COMPLEX, scale, ink, thick)
    return img


def draw_unreadable_form(width=1000, height=1050):
    """A printed table that is not a quiniela slip: find_slip_quad sees it, locate_grid refuses it."""
    img = np.full((height, width, 3), 238, np.uint8)
    for i in range(6):
        cv2.line(img, (100, 200 + i * 120), (900, 200 + i * 120), (30, 30, 30), 4)
    for x in (100, 400, 900):
        cv2.line(img, (x, 200), (x, 800), (30, 30, 30), 4)
    return img


# ----------------------------------------------------------------------------- hands, odd marks and a virtual clock
# (hand guard / mark guard / stack-at-low-fps tests)
SLIP_X = (8, 72, 136, 200)             # draw_slip: the rules around the L / E / V columns ...
SLIP_Y0, SLIP_ROW = 160, 80            # ... the top of the first match row and the row height


def with_hand(frame, x, y, angle=200.0, size=1.7, tone=(110, 150, 205)):
    """A copy of ``frame`` with a skin-coloured hand (see draw_hand) resting at (x, y)."""
    out = frame.copy()
    draw_hand(out, x, y, angle, size, tone)
    return out


def cell_centre(r, c):
    """Centre (x, y) of the L/E/V cell of match row ``r``, column ``c`` in a draw_slip image."""
    return (SLIP_X[c] + SLIP_X[c + 1]) // 2, SLIP_Y0 + r * SLIP_ROW + SLIP_ROW // 2


def _blend(img, draw, alpha):
    if alpha >= 1.0:
        draw(img)
        return img
    ov = img.copy()
    draw(ov)
    img[:] = cv2.addWeighted(ov, alpha, img, 1 - alpha, 0)
    return img


def draw_cross(img, r, c, size=48, thick=4, ink=(170, 60, 25), dx=0, dy=0, alpha=1.0):
    """An X of ``size`` px (a faint one: ``alpha`` < 1 blends the ink into the paper) in cell (r, c), shifted by (dx, dy)."""
    cx, cy = cell_centre(r, c)
    cx, cy, s = cx + dx, cy + dy, size // 2

    def draw(im):
        cv2.line(im, (cx - s, cy - s), (cx + s, cy + s), ink, thick, cv2.LINE_AA)
        cv2.line(im, (cx + s, cy - s), (cx - s, cy + s), ink, thick, cv2.LINE_AA)
    return _blend(img, draw, alpha)


def draw_scribble(img, r, c, strokes=8, thick=5, ink=(170, 60, 25), half=30):
    """A zigzag scribbled over cell (r, c): a crossed-out mark."""
    cx, cy = cell_centre(r, c)
    pts = [(cx - half if i % 2 == 0 else cx + half, cy - half + (2 * half * i) // max(strokes - 1, 1)) for i in range(strokes)]
    for a, b in zip(pts[:-1], pts[1:]):
        cv2.line(img, a, b, ink, thick, cv2.LINE_AA)
    return img


def draw_staple(img, r, c, length=70, angle=20.0, thick=5, grey=(95, 95, 100), dx=0, dy=0):
    """A long thin dark object lying in cell (r, c) at ``angle`` degrees (a staple, a pen slash)."""
    cx, cy = cell_centre(r, c)
    a = np.deg2rad(angle)
    d = (int(round(np.cos(a) * length / 2)), int(round(np.sin(a) * length / 2)))
    cv2.line(img, (cx + dx - d[0], cy + dy - d[1]), (cx + dx + d[0], cy + dy + d[1]), grey, thick, cv2.LINE_AA)
    return img


class VClock:
    """A clock that only moves when told to: ``clock.t += 1 / 30`` per frame makes a test independent of the machine's speed."""

    def __init__(self, t=0.0):
        self.t = t

    def __call__(self):
        return self.t
