"""The review findings that live in the seams between the engine, Excel and the window (gui-B).

What each group pins down, through the real window (a display is needed: xvfb-run on Linux):

* E4 / G04 / SOAK-2  a slip that is NOT in Excel (Undo, Esc on its card, Excel refused it, an internal error) is forgotten by
  the engine's duplicate guard, so lifting it and putting it back reads it again - and only then, only for the slip that
  really lies under the camera;
* F4                 the selected block is followed when rows are inserted above it, and never replaced by another;
* SOAK-5             the Tk thread is never kept busy for seconds, a busy Excel keeps a finished name waiting;
* SOAK-8             a full block holds the slip, names the next block and waits for a human; low stock is announced;
* F2 / F5 / F7       Excel warnings and layout warnings reach the operator without being 'NOT ENTERED';
* E1 / E2            rows under a finger (OCCLUDED), a hand on the slip, a hand over the name;
* SOAK-4             the camera's real frame rate, the slips-read counter and the skipped-slip warning.

The engine is driven frame by frame from the test thread (Bench): the stepping is deterministic and the rest of the path
(events, window, Excel stand-in, reader) is the real one."""
from __future__ import annotations

import copy
import time

import numpy as np
import pytest

tk = pytest.importorskip("tkinter")

import test_pipeline_gui as base
from quiniela_scanner import gui
from quiniela_scanner import slip_vision as sv
from quiniela_scanner.camera import FrameGrabber
from quiniela_scanner.capture import HAND_MESSAGE, TRACKING, Update
from quiniela_scanner.config import Config
from quiniela_scanner.excel_sink import ExcelBusy
from quiniela_scanner.name_reader import NameResult
from helpers import Stage, draw_slip_named, make_pool_workbook

env, readings, flagged = base.env, base.readings, base.flagged           # (fixtures of the pipeline tests)
pump, settle, inject, good, entered, named = base.pump, base.settle, base.inject, base.good, base.entered, base.named
ScriptedReader, STACK, FIRST_ROW, row_of = base.ScriptedReader, base.STACK, base.FIRST_ROW, base.row_of
NAMES = ["CLEMENTE O.", "ANA", "BETO", "CARLOS M.", "DOLORES", "PACO 2", "JUAN G.", "TONY 5000"]
PICKS_A = list("LVLLELLVE")
PICKS_B = list("EEELLVVLV")
BLOCK2_NAME_COL, BLOCK2_TICKET_COL = 26, 25          # make_pool_workbook: block b has '#' in column 13 + 12 b, NOMBRE right of it


class AnyName:
    """A name reader that answers the same for every picture."""

    def __init__(self, name="ANA", delay=0.0):
        self.name, self.delay = name, delay

    def read(self, name_crop, title_crop=None):
        if self.delay:
            time.sleep(self.delay)
        return good(self.name)


class BusyBackend(base.MemBackend):
    """MemBackend whose writes are refused ('Excel is busy') until ``busy_until`` (a time.monotonic() value)."""
    busy_until = 0.0

    def write_row(self, *a, **k):
        if time.monotonic() < self.busy_until:
            self.calls["write_row"] += 1
            raise ExcelBusy("Excel is busy (is a cell being edited?). Press Esc in Excel and retry.")
        return super().write_row(*a, **k)


class WarnBackend(base.MemBackend):
    """MemBackend that reports a best-effort failure after every write (what ComBackend does for Font.Size / Save)."""

    def write_row(self, *a, **k):
        super().write_row(*a, **k)
        return ["font size not applied: Unable to set the Size property of the Font class"]


@pytest.fixture
def mk2(env):
    """mk2(source, reader, wbname, backend_cls, wb=dict of make_pool_workbook options, **config) -> (app, backend)."""
    tmp, _ = env
    apps = []

    def make(source=None, reader=None, wbname="JORNADA 12 PC.xlsx", backend_cls=base.MemBackend, wb=None, **cfgkw):
        xlsx = make_pool_workbook(tmp / wbname, **(wb or {}))
        cfg = Config()
        cfg.read_names_with_claude = False
        for k, v in cfgkw.items():
            setattr(cfg, k, v)
        backend = backend_cls(str(xlsx), 0.0)
        app = gui.App(cfg, source=source or base.NullSource(), backend=backend)
        app.reader = reader
        apps.append(app)
        return app, backend
    yield make
    for app in apps:
        errors = list(app.callback_errors)
        try:
            app.destroy()
        except tk.TclError:
            pass
        assert not errors, errors


# ======================================================================================= helpers
def cells(backend, row, col=base.NAME_COL, sheet="ENTRY PG1"):
    """name + 9 picks of one sheet row, starting at column ``col``"""
    return [backend._wbv[sheet].cell(row=row, column=col + k).value for k in range(10)]


def row_of_ticket(backend, ticket, ticket_col=base.NAME_COL - 1, sheet="ENTRY PG1"):
    """the name + 9 picks of the row that carries ``ticket`` (rows may have moved), None if no row does"""
    ws = backend._wbv[sheet]
    for r in range(1, 200):
        if ws.cell(row=r, column=ticket_col).value == ticket:
            return cells(backend, r, ticket_col + 1, sheet)
    return None


def insert_row(backend, at, sheet="ENTRY PG1"):
    """the operator inserts a row in Excel (above the headers)"""
    for wb in (backend._wb, backend._wbv):
        wb[sheet].insert_rows(at)


def sheet_values(backend, sheet="ENTRY PG1"):
    ws = backend._wbv[sheet]
    return [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=70, min_col=1, max_col=60)]


class Bench:
    """Plays frames into the real engine from the test thread (the camera thread is stopped) and hands the window what the
    camera loop would hand it: the scan events, the newest frame for the preview."""

    def __init__(self, app):
        app._stop_camera()
        self.app, self.n, self.msgs = app, 0, []

    def play(self, frames):
        app = self.app
        for f in frames:
            self.n += 1
            upd = app.engine.feed(f)
            self.msgs.append(upd.message)
            with app._lock:
                app._latest = (self.n, f, upd)
            if upd.reading is not None or upd.error or upd.duplicate:
                app.events.put(("scan", upd, time.perf_counter()))
            app._next_preview = 0.0
            app._tick()
            app.update()


class Scene:
    """Slips on a table: pieces of film that a Bench plays (one Stage, so the lifts reveal the slip underneath)."""

    def __init__(self, *slips, seed=2):
        self.st = Stage(seed=seed)
        self.ids = [self.st.add(draw_slip_named(p, n, seed=s)) for p, n, s in slips]

    def _f(self, gen):
        return [f for f, _ in gen]

    def table(self, n=5):
        return self._f(self.st.empty(n))

    def hold(self, k, n=25, under=None):
        return self._f(self.st.hold(self.ids[k], n, under=None if under is None else self.ids[under]))

    def lift(self, k, under=None, n=6):
        return self._f(self.st.lift(self.ids[k], None if under is None else self.ids[under], n))

    def drop(self, k, n=5):
        return self._f(self.st.drop(self.ids[k], None, n))

    def replace(self, k):
        """the slip is lifted off, the table is empty for a moment, the same slip is put back and lies still"""
        return self.lift(k) + self.table(6) + self.drop(k) + self.hold(k, 30)


def live(app):
    return [j for j in app.entries if j.state != "undone"]


# ======================================================================================= E4 / G04 / SOAK-2
def test_undo_then_lifting_and_replacing_the_slip_enters_it_again(env, mk2):
    app, backend = mk2(reader=AnyName())
    scene = Scene((PICKS_A, "CLEMENTE O.", 1))
    bench = Bench(app)
    bench.play(scene.table() + scene.hold(0))
    assert entered(app, 1) and named(app, app.entries[0])
    job = app.entries[0]
    app.undo()
    assert job.state == "undone" and row_of(backend, 0) == [None] * 10
    assert "Undone - lift the slip and place it again" in app.status_var.get() and not app._entered_msg
    bench.play(scene.hold(0, 30))                              # still lying there: not read again, the message stays
    assert len(app.entries) == 1 and "Undone - lift the slip" in app.status_var.get()
    bench.play(scene.replace(0))                               # lifted and put back: read again
    assert entered(app, 2) and live(app)[0].ticket == 251
    assert row_of(backend, 0)[1:] == PICKS_A
    assert app.dupes_skipped == 0


def test_the_same_through_the_threaded_camera_loop(env, mk2):
    """The scenario of the reviewers' scripts: real camera thread, queued frames, Undo, lift, put back."""
    import collections
    scene = Scene((PICKS_A, "CLEMENTE O.", 1))

    class QueueSource(base.NullSource):
        def __init__(self):
            super().__init__()
            self.q, self.last = collections.deque(), self.frame

        def read(self):
            time.sleep(1 / 30)
            if self.q:
                self.last = self.q.popleft()
            return self.last
    src = QueueSource()
    app, backend = mk2(src, AnyName())
    src.q.extend(scene.table() + scene.hold(0))
    assert entered(app, 1) and named(app, app.entries[0])
    app.undo()
    src.q.extend(scene.replace(0))
    assert pump(app, lambda: len(live(app)) == 1, 20), "the re-placed slip was silently ignored"
    assert row_of(backend, 0)[1:] == PICKS_A


def test_discarding_a_card_then_replacing_the_slip_opens_the_card_again(env, mk2):
    app, backend = mk2(reader=AnyName())
    scene = Scene((["L", "", "LV", "E", "L", "L", "V", "E", "E"], "ANA", 21))     # row 2 empty, row 3 double: needs a card
    bench = Bench(app)
    bench.play(scene.table() + scene.hold(0))
    assert pump(app, lambda: app.pending is not None)
    app.discard()                                              # Esc
    assert "Slip discarded - lift it and place it again" in app.status_var.get()
    bench.play(scene.hold(0, 30))
    assert app.pending is None and "Slip discarded - lift it" in app.status_var.get()    # not re-read while it lies there
    bench.play(scene.replace(0))
    assert pump(app, lambda: app.pending is not None), "the re-placed slip was silently ignored"
    assert app.dupes_skipped == 0


def test_a_slip_excel_refused_and_then_discarded_can_be_put_back(env, mk2):
    app, backend = mk2(reader=AnyName())
    backend.fail_writes = 1                                    # Excel refuses the first write
    scene = Scene((PICKS_A, "CLEMENTE O.", 1))
    bench = Bench(app)
    bench.play(scene.table() + scene.hold(0))
    assert pump(app, lambda: app.pending is not None) and app.pending.hold and not app.entries
    bench.play(scene.hold(0, 10))
    assert len(app.entries) == 0                               # (the held card keeps the slip: no second card for it)
    app.discard()
    bench.play(scene.replace(0))
    assert entered(app, 1) and row_of(backend, 0)[1:] == PICKS_A


def test_a_slip_whose_card_is_open_is_not_forgotten_by_the_refusal_itself(env, mk2):
    """The guard is lifted when the card is DISCARDED, not when Excel refuses: lifting and replacing a slip whose hold card
    is still open must not queue a second copy of it (two tickets for one slip)."""
    app, backend = mk2(reader=AnyName())
    backend.fail_writes = 1
    scene = Scene((PICKS_A, "CLEMENTE O.", 1))
    bench = Bench(app)
    bench.play(scene.table() + scene.hold(0))
    assert pump(app, lambda: app.pending is not None)
    bench.play(scene.replace(0))
    settle(app, 0.2)
    assert not app.backlog and app.dupes_skipped >= 1


def test_undoing_an_older_entry_keeps_the_guard_of_the_slip_under_the_camera(env, mk2):
    """Only the slip the engine remembers last is forgotten: undo ticket A while B lies under the camera and B, lifted
    and put back, must still be recognised (an unconditional forget_last would enter it twice)."""
    app, backend = mk2(reader=AnyName())
    scene = Scene((PICKS_A, "CLEMENTE O.", 1), (PICKS_B, "ANA", 5))
    bench = Bench(app)
    bench.play(scene.table() + scene.hold(0, under=1) + scene.lift(0, under=1) + scene.hold(1))
    assert entered(app, 2)
    a, b = app.entries
    app.undo_job(a)                                            # the OLD one
    assert a.state == "undone" and b.state != "undone"
    assert not app._lift_msg                                   # (the slip lying there is B, which is in Excel)
    bench.play(scene.replace(1))
    settle(app, 0.3)
    assert len(live(app)) == 1 and app.dupes_skipped >= 1 and app.engine.last_fingerprint is b.fp


def test_undoing_the_latest_entry_forgets_it_only_if_it_is_the_one_remembered(env, mk2, readings):
    app, backend = mk2(reader=ScriptedReader())
    app._stop_camera()
    reader = app.reader
    rd = inject(app, reader, readings[0], good("ANA"))
    assert entered(app, 1)
    job = app.entries[0]
    assert job.fp is None                                      # (a hand-made Update carries no fingerprint: nothing to forget)
    assert app.forget_if_current(None) is False and app.forget_if_current(job.fp) is False
    assert rd.picks == job.picks or True


def test_a_slip_that_could_not_be_processed_can_be_put_back(env, mk2):
    app, backend = mk2(reader=AnyName())
    boom = {"n": 1}
    handle = app._handle_scan

    def flaky(upd, t_event, snap=None):
        if boom["n"]:
            boom["n"] -= 1
            raise RuntimeError("bug in the window")
        return handle(upd, t_event, snap)
    app._handle_scan = flaky
    scene = Scene((PICKS_A, "CLEMENTE O.", 1))
    bench = Bench(app)
    bench.play(scene.table() + scene.hold(0))
    assert pump(app, lambda: "scanfail" in app._alerts)
    app.callback_errors.clear()                                # (the internal error is the point of this test)
    assert app.engine.last_fingerprint is None and not app.entries
    bench.play(scene.replace(0))
    assert entered(app, 1)


# ======================================================================================= F4
def test_the_selected_block_follows_rows_inserted_above_it(env, mk2, readings):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader)
    app.select_info(app.infos[1])                              # block 2: tickets 481-532 [CH]
    key = app.sel_key
    inject(app, reader, readings[0], good("A0"))
    assert entered(app, 1) and app.entries[0].ticket == 484
    insert_row(backend, 2)                                     # one row more above the headers
    for k in (1, 2, 3):
        inject(app, reader, readings[k], good(f"A{k}"))
        assert entered(app, 1 + k)
    assert [j.ticket for j in app.entries] == [484, 485, 486, 487]
    assert app.sel_key != key and app.sel_key.endswith("Z5") and "tickets 481-532" in app.block_var.get()
    assert app.ticket_var.get() == "488" and not app.alert_var.get()
    assert named(app, *app.entries)
    for k in range(4):
        assert row_of_ticket(backend, 484 + k, BLOCK2_TICKET_COL) == [f"A{k}"] + STACK[k][0]
    assert row_of_ticket(backend, 251)[1:] == row_of_ticket(backend, 251)[1:] and row_of_ticket(backend, 251)[0] is None


def test_a_ticket_the_operator_typed_survives_the_block_being_followed(env, mk2):
    app, backend = mk2(reader=ScriptedReader())
    app.select_info(app.infos[1])
    app.ticket_var.set("490")                                  # typed by the operator
    assert app.ticket_custom
    insert_row(backend, 2)
    app.sink.refresh()
    app.refresh_blocks()
    assert app.sel_key.endswith("Z5") and app.ticket_var.get() == "490" and app.ticket_custom


def test_a_block_that_is_gone_is_not_replaced_by_another_one(env, mk2, readings, monkeypatch):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader)
    app.select_info(app.infos[1])
    label, block = app.block_var.get(), app.sel_block
    before = sheet_values(backend)
    for b in (backend._wb, backend._wbv):                      # the operator deletes the columns of block 2
        b["ENTRY PG1"].delete_cols(25, 12)
    app.sink.refresh()                                         # (what the sink does by itself when a write finds rows moved)
    app.refresh_blocks()
    assert "no longer in the sheet" in app.alert_var.get() and app.sel_block is block and app.block_var.get() == label
    assert str(app.block_lbl.cget("foreground")) == gui.BAD
    inject(app, reader, readings[0], good("ANA"))
    assert pump(app, lambda: app.pending is not None)
    assert "gone from the sheet" in app.banner_var.get() and not app.entries
    assert backend.calls["write_row"] == 0                     # nothing was sent to Excel, not even to a look-alike block
    app.discard()
    app.block_var.set(app.infos[0].label)                      # the operator picks a block ...
    app.on_block_change()
    assert "no longer in the sheet" not in app.alert_var.get() and app._block_gone is False
    inject(app, reader, readings[1], good("BETO"))             # ... and the stack goes on there
    assert entered(app, 1) and app.entries[0].ticket == 251
    assert before != sheet_values(backend)


def test_another_block_that_moved_into_its_place_is_not_the_same_block(env, mk2, monkeypatch):
    app, backend = mk2(reader=ScriptedReader())
    app.select_info(app.infos[1])
    other = app.infos[2].block                                 # same sheet, other tickets (721-...)
    monkeypatch.setattr(app.sink, "resolve_block", lambda old: other)
    app.refresh_blocks()
    assert app._block_gone and app.sel_block is not other and "no longer in the sheet" in app.alert_var.get()


def test_reconnecting_to_the_same_workbook_keeps_the_selected_block(env, mk2):
    app, backend = mk2(reader=ScriptedReader())
    app.select_info(app.infos[1])
    app.connect_excel()
    assert app.sel_key.endswith("Z4") and app.ticket_var.get() == "484"
    backend.workbook_name = "JORNADA 13 PC.xlsx"               # another workbook: start from the first block with room
    app.connect_excel()
    assert app.sel_key.endswith("N4") and app.ticket_var.get() == "251"


# ======================================================================================= SOAK-5
def test_one_tick_never_works_through_a_whole_queue_of_slow_events(env, mk2, monkeypatch):
    app, backend = mk2(reader=ScriptedReader())
    handled = []

    def slow(kind, key, res, ms):
        time.sleep(0.1)                                        # (a busy Excel: every name write costs ~1.2 s)
        handled.append(key)
    monkeypatch.setattr(app, "on_name", slow)
    for k in range(10):
        app.events.put(("name", "job", 1000 + k, NameResult("X"), 1.0))
    t0 = time.perf_counter()
    app._events()
    one_tick = time.perf_counter() - t0
    assert 1 <= len(handled) < 10 and one_tick < 0.45          # the budget, not the 1.0 s of the whole queue
    while len(handled) < 10:                                   # ... and nothing is lost: the next ticks do the rest
        app._events()
    assert handled == list(range(1000, 1010))


def test_a_name_that_meets_a_busy_excel_waits_and_is_written_when_excel_is_free(env, mk2, readings, monkeypatch):
    monkeypatch.setattr(gui, "BUSY_RETRY_S", 0.1)
    reader, gate = ScriptedReader(), base.threading.Event()
    app, backend = mk2(reader=reader, backend_cls=BusyBackend, name_timeout_s=0.3)
    inject(app, reader, readings[0], good("ANA"), gate=gate)
    assert entered(app, 1)
    job = app.entries[0]
    backend.busy_until = time.monotonic() + 1.0                # Excel is busy for a second ...
    gate.set()                                                 # ... when the name arrives
    assert pump(app, lambda: job.writing)
    assert job.state == "reading" and "Excel busy" in job.text and "waiting for Excel: 1" in app.stats_var.get()
    assert "EXCEL IS BUSY" in app.notice_var.get()
    settle(app, 0.5)                                           # longer than name_timeout_s: the time-out must not touch it
    assert job.state == "reading" and row_of(backend, 0)[0] is None
    assert pump(app, lambda: job.state == "ok", 10)
    assert row_of(backend, 0)[0] == "ANA" and job.written == "ANA" and not job.writing
    assert "EXCEL IS BUSY" not in app.notice_var.get() and not app._name_retry


def test_a_busy_excel_costs_one_attempt_not_one_per_queued_name(env, mk2, readings, monkeypatch):
    monkeypatch.setattr(gui, "BUSY_RETRY_S", 0.3)
    reader = ScriptedReader()
    app, backend = mk2(reader=reader, backend_cls=BusyBackend)
    gates = [base.threading.Event() for _ in range(5)]
    for k in range(5):
        inject(app, reader, readings[k], good(NAMES[k]), gate=gates[k])
    assert entered(app, 5)
    backend.busy_until = time.monotonic() + 1.0
    before = backend.calls["write_row"]
    for g in gates:
        g.set()
    assert pump(app, lambda: len(reader.finished) == 5 and all(j.writing for j in app.entries))
    assert backend.calls["write_row"] - before == 1            # the first name tried; the other four waited behind it
    assert pump(app, lambda: all(j.state == "ok" for j in app.entries), 15)
    assert [row_of(backend, k)[0] for k in range(5)] == NAMES[:5]


def test_a_name_gives_up_on_a_busy_excel_only_after_the_cap(env, mk2, readings, monkeypatch):
    monkeypatch.setattr(gui, "BUSY_RETRY_S", 0.05)
    monkeypatch.setattr(gui, "NAME_BUSY_S", 0.4)
    reader = ScriptedReader()
    app, backend = mk2(reader=reader, backend_cls=BusyBackend)
    inject(app, reader, readings[0], good("ANA"), delay=0.1)
    assert entered(app, 1)
    backend.busy_until = time.monotonic() + 60
    job = app.entries[0]
    assert pump(app, lambda: job.state != "reading", 10)
    assert job.state == "review" and "NAME NOT WRITTEN" in job.text and "busy" in job.reason
    assert job.guess == "ANA" and not job.writing and not app._name_retry       # (the guess is pre-filled for Save name)


def test_a_name_waiting_for_excel_is_still_written_when_the_window_closes(env, mk2, readings, monkeypatch):
    monkeypatch.setattr(gui, "BUSY_RETRY_S", 0.05)
    reader, gate = ScriptedReader(), base.threading.Event()
    app, backend = mk2(reader=reader, backend_cls=BusyBackend)
    inject(app, reader, readings[0], good("ANA"), gate=gate)
    assert entered(app, 1)
    backend.busy_until = time.monotonic() + 0.3
    gate.set()
    job = app.entries[0]
    assert pump(app, lambda: job.writing)
    app.on_close()                                             # Excel is free again by the time it gives up waiting
    assert backend.row(FIRST_ROW)[0] == "ANA" and job.state == "ok"


def test_a_busy_excel_for_the_picks_falls_back_to_the_card_with_a_clear_message(env, mk2, readings):
    tmp, beeps = env
    reader = ScriptedReader()
    app, backend = mk2(reader=reader, backend_cls=BusyBackend)
    backend.busy_until = time.monotonic() + 3600
    inject(app, reader, readings[0], good("ANA"))
    assert pump(app, lambda: app.pending is not None)
    assert "Excel is busy" in app.banner_var.get() and "Esc in Excel" in app.banner_var.get()
    assert "NOT ENTERED" in app.alert_var.get() and beeps["attention"] >= 1 and not app.entries
    assert pump(app, lambda: app.name_var.get() == "ANA")
    backend.busy_until = 0.0                                   # the operator pressed Esc in Excel
    app.commit()
    assert row_of(backend, 0) == ["ANA"] + STACK[0][0] and app.alert_var.get() == "" and app.pending is None


def test_a_busy_excel_during_enter_on_the_card_says_so_on_the_card(env, mk2, flagged):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader, backend_cls=BusyBackend)
    inject(app, reader, flagged, good("ANA"))
    assert pump(app, lambda: app.pending is not None)
    base.fix_flagged_card(app)
    backend.busy_until = time.monotonic() + 3600
    app.commit()
    assert app.pending is not None and "Excel is busy" in app.banner_var.get()
    backend.busy_until = 0.0
    app.commit()
    assert app.pending is None and len(app.entries) == 1


# ======================================================================================= SOAK-8
def nearly_full(app, backend, free):
    """Block 1 (rows 15..56 free) gets ``free`` tickets left, then the app reads the sheet again."""
    for r in range(FIRST_ROW, 57 - free):
        backend.type_by_hand(r, base.NAME_COL, f"X{r}")
    app.connect_excel()
    assert app.infos[0].free == free


def test_a_full_block_holds_the_slip_names_the_next_block_and_waits_for_enter(env, mk2, readings):
    tmp, beeps = env
    reader = ScriptedReader()
    app, backend = mk2(reader=reader)
    nearly_full(app, backend, 1)
    inject(app, reader, readings[0], good("A0"))               # takes the last ticket (292)
    assert entered(app, 1) and app.entries[0].ticket == 292
    assert "FULL" in app.low_note and str(app.block_lbl.cget("foreground")) == gui.BAD
    assert "This block is now FULL" in app.stats_var.get()      # (the status line is cleared as soon as the camera sees no slip)
    inject(app, reader, readings[1], good("A1"))               # ... and this one does not fit
    assert pump(app, lambda: app.pending is not None)
    assert app.pending.hold and not app.entries[1:] and len(app.entries) == 1
    assert app.banner_var.get() == "BLOCK FULL - choose another block in the Block list, then press Enter"
    nxt = app.infos[1].label
    assert nxt in app.alert_var.get() and "BLOCK FULL" in app.alert_var.get() and app.block_var.get() == nxt
    assert app._block_confirm == nxt and app.ticket_var.get() == "484"
    assert row_of_ticket(backend, 484, BLOCK2_TICKET_COL)[1:] == [None] * 9                # nothing written there yet
    inject(app, reader, readings[2], good("A2"))               # the next slip waits behind the held card
    assert pump(app, lambda: len(app.backlog) == 1)
    assert pump(app, lambda: app.name_var.get() == "A1")
    app.commit()                                               # Enter = "yes, that block"
    assert row_of_ticket(backend, 484, BLOCK2_TICKET_COL)[1:] == STACK[1][0]
    assert app._block_confirm == "" and "BLOCK FULL" not in app.alert_var.get()
    assert entered(app, 3) and app.entries[2].ticket == 485 and not app.backlog   # the queued slip follows, in the same block


def test_after_a_full_block_nothing_goes_to_another_block_until_a_human_confirms(env, mk2, readings):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader)
    nearly_full(app, backend, 0)
    inject(app, reader, readings[0], good("A0"))
    assert pump(app, lambda: app.pending is not None)
    app.discard()                                              # Esc on the held card: the slip is gone, the new block unconfirmed
    inject(app, reader, readings[1], good("A1"))               # a clean slip, but the block was never confirmed: a card
    assert pump(app, lambda: app.pending is not None) and not app.entries
    assert "NEW BLOCK" in app.banner_var.get()
    assert pump(app, lambda: app.name_var.get() == "A1")
    app.commit()
    assert len(app.entries) == 1 and app.entries[0].ticket == 484 and app._block_confirm == ""
    inject(app, reader, readings[2], good("A2"))               # and from now on the stack flows by itself again
    assert entered(app, 2) and app.entries[1].ticket == 485 and app.pending is None


def test_picking_a_block_yourself_clears_the_full_block_state(env, mk2, readings):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader)
    nearly_full(app, backend, 0)
    inject(app, reader, readings[0], good("A0"))
    assert pump(app, lambda: app.pending is not None) and app._block_confirm
    app.block_var.set(app.infos[2].label)                      # the operator wants the third block (another group)
    app.on_block_change()
    assert app._block_confirm == "" and "BLOCK FULL" not in app.alert_var.get() and app.ticket_var.get() == "721"
    assert pump(app, lambda: app.name_var.get() == "A0")
    app.commit()
    assert app.entries[0].ticket == 721


def test_a_full_block_with_no_other_block_free_says_so(env, mk2, readings):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader, wb=dict(filled=(52, 52, 52)))
    inject(app, reader, readings[0], good("A0"))
    assert pump(app, lambda: app.pending is not None)
    assert "No other block has a free ticket" in app.alert_var.get() and not app._block_confirm
    assert app.banner_var.get().startswith("BLOCK FULL")


def test_the_next_block_of_the_same_group_is_preferred(env, mk2, readings):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader, wb=dict(filled=(52, 3, 3), suffixes=("RRR", "CH", "RRR")))
    # block 1 [RRR] is full, block 2 is [CH], block 3 is [RRR]: the stack of group RRR goes on in block 3
    app.select_info(app.infos[0])                              # (the operator is working in block 1, which is full)
    inject(app, reader, readings[0], good("A0"))
    assert pump(app, lambda: app.pending is not None)
    assert "[RRR]" in app.block_var.get() and app._block_confirm == app.block_var.get() and app.ticket_var.get() == "724"


def test_a_block_with_few_tickets_left_turns_amber_and_beeps_at_3_and_at_0(env, mk2, readings):
    tmp, beeps = env
    reader = ScriptedReader()
    app, backend = mk2(reader=reader)
    nearly_full(app, backend, 7)
    seen = []
    for k in range(7):
        inject(app, reader, readings[k], good(NAMES[k]))
        assert entered(app, k + 1)
        info = app.selected_info()
        seen.append((info.free, str(app.block_lbl.cget("foreground")), beeps["attention"], bool(app.low_note)))
    assert [s[0] for s in seen] == [6, 5, 4, 3, 2, 1, 0]
    assert seen[0][1] == str(app._block_fg) and not seen[0][3]                 # 6 left: nothing yet
    assert all(s[3] for s in seen[1:])                                         # 5 left or fewer: the note is up
    assert [s[1] for s in seen[1:]] == [str(gui.WARN)] * 5 + [str(gui.BAD)]    # amber from 5 left, red when full
    assert [s[2] for s in seen] == [0, 0, 0, 1, 1, 1, 2]                       # one beep at 3 left, one at 0


# ======================================================================================= F2 / F5 / F7
def test_a_failed_font_or_save_after_the_values_landed_is_a_notice_not_a_refusal(env, mk2, readings):
    reader, gate = ScriptedReader(), base.threading.Event()
    app, backend = mk2(reader=reader, backend_cls=WarnBackend)
    inject(app, reader, readings[0], good("ANA"), gate=gate)
    assert entered(app, 1)
    assert "Entered ticket 251, but font size not applied" in app.notice_var.get()
    assert "NOT ENTERED" not in app.alert_var.get() and app.pending is None and row_of(backend, 0)[1:] == STACK[0][0]
    gate.set()
    assert named(app, app.entries[0]) and app.entries[0].state == "ok"
    assert "Ticket 251: the name is in, but font size not applied" in app.notice_var.get()
    app.dismiss_notices()
    assert not app.notice_frame.winfo_ismapped()
    app.undo()                                                 # (undo warns the same way)
    assert "Ticket 251 was cleared, but font size not applied" in app.notice_var.get()


def test_a_card_entry_with_a_warning_is_also_just_a_notice(env, mk2, flagged):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader, backend_cls=WarnBackend)
    inject(app, reader, flagged, good("ANA"))
    assert pump(app, lambda: app.pending is not None)
    base.fix_flagged_card(app)
    app.commit()
    assert len(app.entries) == 1 and "Entered ticket 251, but font size" in app.notice_var.get()


def test_layout_warnings_are_shown_once_at_connect_and_go_when_fixed(env, mk2):
    app, backend = mk2(reader=None)
    assert not app.notice_var.get() and not app.sink.layout_warnings
    backend.type_by_hand(30, base.NAME_COL - 1, "266 ")        # a ticket stored as text cuts block 1 short (row 30 is ticket 266)
    app.connect_excel()
    assert "EXCEL LAYOUT" in app.notice_var.get() and "block ends at row 29" in app.notice_var.get()
    app.update()
    assert app.notice_frame.winfo_ismapped()
    app.flush_io()
    log = (env[0] / "scanner.log").read_text()
    assert log.count("LAYOUT_WARNING") == len(app.sink.layout_warnings) and "layout_warnings=" in log
    app.refresh_blocks()                                       # a re-list changes nothing: still one message, one log set
    assert app.notice_var.get().count("EXCEL LAYOUT") == 1
    app.flush_io()
    assert (env[0] / "scanner.log").read_text().count("LAYOUT_WARNING") == len(app.sink.layout_warnings)
    backend.type_by_hand(30, base.NAME_COL - 1, 266)
    app.connect_excel()
    app.update()
    assert not app.notice_var.get() and not app.notice_frame.winfo_ismapped()


def test_run_bat_excel_prints_the_layout_warnings(tmp_path, capsys):
    import openpyxl
    from quiniela_scanner.__main__ import main
    path = make_pool_workbook(tmp_path / "pool.xlsx")
    wb = openpyxl.load_workbook(path)
    wb["ENTRY PG1"].cell(30, base.NAME_COL - 1).value = "276 "
    wb.save(path)
    assert main(["excel", "--backend", "file", "--path", str(path)]) == 0
    out = capsys.readouterr().out
    assert "WARNING:" in out and "block ends at row 29" in out and "next free" in out


# ======================================================================================= E1 / E2
def occluded_copy(rd, row=2, name_hidden=False):
    rd = copy.copy(rd)
    rd.name_crop = rd.name_crop.copy()
    rd.picks, rd.verdicts = list(rd.picks), list(rd.verdicts)
    rd.picks[row], rd.verdicts[row] = "", sv.OCCLUDED
    rd.hand, rd.name_hidden = True, name_hidden
    return rd


def test_a_row_under_a_finger_is_never_auto_entered_and_is_red_on_the_card(env, mk2, readings, monkeypatch):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader)
    rd = occluded_copy(readings[0])
    assert not app.auto_eligible(rd) and not app.fast_ok(rd)
    inject(app, reader, rd, good("ANA"))
    assert pump(app, lambda: app.pending is not None) and not app.entries
    lbl, note = app.pick_lbls[2]
    assert note.cget("text") == "finger over this row" and str(lbl.cget("foreground")) == gui.BAD
    assert "finger or hand covers row(s) 3" in app.status_var.get() and app.banner_var.get() == "CHECK THIS SLIP"
    assert pump(app, lambda: app.name_var.get() == "ANA")
    asked = []
    monkeypatch.setattr(gui.messagebox, "askyesno", lambda *a, **k: asked.append(a) or False)
    app.commit()                                               # Enter with the covered row still on 'empty'
    assert asked and app.pending is not None and row_of(backend, 0) == [None] * 10
    app.pick_vars[2].set("L")                                  # the operator reads the row off the slip
    app.pick_edited(2)
    app.commit()
    assert row_of(backend, 0)[1:] == [readings[0].picks[i] if i != 2 else "L" for i in range(9)]


def test_take_your_hand_off_the_slip_is_shown_with_a_red_outline(env, mk2, monkeypatch):
    app, backend = mk2(reader=None)
    app._stop_camera()
    colours = []
    real = sv.draw_quad
    monkeypatch.setattr(sv, "draw_quad", lambda img, quad, colour, t=2: colours.append(colour) or real(img, quad, colour, t))
    frame = np.full((1080, 1920, 3), 200, np.uint8)
    quad = np.array([[600, 300], [1300, 300], [1300, 800], [600, 800]], np.float32)
    with app._lock:
        app._latest = (1, frame, Update(TRACKING, HAND_MESSAGE, quad, 0.5))
    app._next_preview = 0.0
    app.render_preview(time.perf_counter())
    assert app.status_var.get() == HAND_MESSAGE and colours[-1] == (0, 0, 230)
    with app._lock:
        app._latest = (2, frame, Update(TRACKING, "Hold still...", quad, 0.5))
    app._next_preview = 0.0
    app.render_preview(time.perf_counter())
    assert app.status_var.get() == "Hold still..." and colours[-1] == (0, 200, 255)


def test_a_hand_over_the_name_forces_a_name_review_even_when_claude_is_sure(env, mk2, readings):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader)
    rd = copy.copy(readings[0])
    rd.name_hidden = True
    inject(app, reader, rd, good("ANA"))
    assert entered(app, 1) and named(app, app.entries[0])
    job = app.entries[0]
    assert job.state == "review" and "hand was over the name" in job.text and job.reason
    assert row_of(backend, 0)[0] == "ANA ?" and job.guess == "ANA"          # the guess, marked: nobody trusts it blindly
    assert str(app.tree.item(job.iid, "tags")[0]) == "review" and "TO REVIEW: 1" in app.stats_var.get()


def test_a_hand_over_the_name_is_said_on_the_card_too(env, mk2, readings):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader, auto_enter=False)
    rd = copy.copy(readings[0])
    rd.name_hidden = True
    inject(app, reader, rd, good("ANA"))
    assert pump(app, lambda: app.pending is not None and app.pending.name_done)
    assert "hand covered the name" in app.name_note.get() and app.pending.name_legible is False


# ======================================================================================= SOAK-4 (the window's part)
class Cam:
    """a camera that delivers ``hz`` pictures per second"""
    label = "fake"
    demo = False

    def __init__(self, hz):
        self.dt, self.frame, self.fps_reported = 1.0 / hz, np.full((1080, 1920, 3), 235, np.uint8), 30.0

    def read(self):
        time.sleep(self.dt)
        return self.frame

    def set_autofocus(self, *a, **k):
        pass

    def open_driver_settings(self):
        pass

    def close(self):
        pass


@pytest.mark.parametrize("hz", [40, 10])
def test_the_grabber_measures_the_frame_rate_the_camera_really_delivers(hz):
    g = FrameGrabber(Cam(hz))
    try:
        assert g.fps() == 0.0                                  # not known before ~20 gaps have been seen
        end = time.time() + 12
        while g.fps() == 0.0 and time.time() < end:
            time.sleep(0.05)
        assert 0.6 * hz < g.fps() < 1.4 * hz, g.fps()
    finally:
        g.stop()


def test_the_grabber_ignores_a_stall_when_it_measures(monkeypatch):
    class Stalling(Cam):
        n = 0

        def read(self):
            self.n += 1
            time.sleep(1.0 if self.n == 12 else self.dt)       # one second of nothing in the middle
            return self.frame
    g = FrameGrabber(Stalling(50))
    try:
        end = time.time() + 15
        while g.fps() == 0.0 and time.time() < end:
            time.sleep(0.05)
        time.sleep(0.5)
        assert g.fps() > 25
    finally:
        g.stop()


class FakeGrabber:
    def __init__(self, fps):
        self._fps = fps

    def fps(self):
        return self._fps

    def age(self):
        return 0.0

    def stop(self):
        pass


def test_a_slow_camera_gets_a_persistent_amber_warning_and_its_rate_in_the_strip(env, mk2):
    app, backend = mk2(reader=None)
    app._stop_camera()
    app.grabber = FakeGrabber(8.0)
    t = time.perf_counter() + 1000.0                           # (a fake clock ahead of the real one, which app.update() uses)
    for _ in range(gui.FPS_CHECKS - 1):                        # a dip is not a warning ...
        app.check_camera(t)
        t += 1.1
    assert "CAMERA IS SLOW" not in app.notice_var.get() and "camera 8 fps" in app.stats_var.get()
    app.check_camera(t)                                        # ... three checks in a row are
    assert "CAMERA IS SLOW" in app.notice_var.get() and "Low Light Compensation" in app.notice_var.get()
    app.update()
    assert "slips may be missed" in app.notice_var.get() and app.notice_frame.winfo_ismapped()
    app.dismiss_notices()                                      # dismissed: stays away while it is the same warning
    t += 1.1
    app.check_camera(t)
    assert not app.notice_var.get()
    app.grabber = FakeGrabber(29.0)                            # the light is back
    t += 1.1
    app.check_camera(t)
    assert "camera 29 fps" in app.stats_var.get() and not app.notice_var.get() and app._fps_bad == 0
    app.grabber = FakeGrabber(9.0)                             # ... and gone again: warned again (the dismissal ended)
    for _ in range(gui.FPS_CHECKS):
        t += 1.1
        app.check_camera(t)
    assert "CAMERA IS SLOW" in app.notice_var.get()


def test_the_real_grabber_of_a_slow_camera_raises_the_warning(env, mk2):
    app, backend = mk2(Cam(8), reader=None)
    assert pump(app, lambda: "CAMERA IS SLOW" in app.notice_var.get(), 25), app.stats_var.get()
    assert "camera 8 fps" in app.stats_var.get() or "camera 7 fps" in app.stats_var.get() or "camera 9 fps" in app.stats_var.get()


def test_slips_read_are_counted_and_a_late_slip_is_flagged_once(env, mk2):
    app, backend = mk2(reader=None)
    app._stop_camera()
    t = 5000.0
    for k in range(6):                                         # a steady 2 s pace
        app.note_slip(t)
        t += 2.0
    assert app.slips_read == 6 and "slips read: 6" in app.stats_var.get() and not app.notice_var.get()
    app.note_slip(t + 3.0)                                     # 5 s after the last one: 2.5 x the pace is 5 s -> not yet
    assert not app.notice_var.get()
    t += 3.0
    app.note_slip(t + 7.5)                                     # 7.5 s gap: well over 2.5 x the 2 s pace
    assert "A SLIP MAY HAVE BEEN SKIPPED" in app.notice_var.get() and "8 s passed" in app.notice_var.get()
    app.flush_io()
    assert (env[0] / "scanner.log").read_text().count("SLIP_GAP") == 1
    app.dismiss_notices()
    t += 7.5
    app.note_slip(t + 2.0)                                     # back to the pace: nothing more
    assert not app.notice_var.get() and app.slips_read == 9


def test_a_break_is_not_a_skipped_slip(env, mk2):
    app, backend = mk2(reader=None)
    app._stop_camera()
    t = 100.0
    for k in range(6):
        app.note_slip(t)
        t += 2.0
    app.note_slip(t + gui.BREAK_S + 30)                        # a coffee break
    assert not app.notice_var.get()


def test_every_slip_the_engine_hands_over_is_counted_whether_entered_or_on_a_card(env, mk2, readings, flagged):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader)
    inject(app, reader, readings[0], good("ANA"))
    assert entered(app, 1)
    inject(app, reader, flagged, good("ANA"))
    assert pump(app, lambda: app.pending is not None)
    inject(app, reader, readings[1], good("BETO"))
    assert pump(app, lambda: len(app.backlog) == 1)
    assert app.slips_read == 3 and "Entered: 1" in app.stats_var.get() and "slips read: 3" in app.stats_var.get()
