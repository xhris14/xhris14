import openpyxl
import pytest

from quiniela_scanner.excel_sink import (ComBackend, ExcelError, ExcelSink, OpenpyxlBackend, SlotTaken,
                                         discover_blocks, needs_text_prefix, Snapshot)
from helpers import make_pool_workbook


@pytest.fixture
def sink(tmp_path):
    path = make_pool_workbook(tmp_path / "JORNADA 10 PC.xlsx")
    return ExcelSink(OpenpyxlBackend(str(path))), path


def test_discovers_all_three_blocks(sink):
    s, _ = sink
    infos = s.block_infos()
    assert [i.block.name_col for i in infos] == [14, 26, 38]        # N, Z, AL
    assert [i.free for i in infos] == [42, 49, 52]
    assert [i.suffix for i in infos] == ["RRR", "CH", ""]
    assert infos[0].first_free.ticket == 251 and infos[1].first_free.ticket == 484
    assert infos[0].block.ticket_col == 13
    assert "tickets 241-292" in infos[0].label


def test_next_slot_and_default_block(sink):
    s, _ = sink
    assert s.default_block().first_free.ticket == 251
    assert s.next_slot().ticket == 251
    assert s.slot_for_ticket(493).row == 5 + 12


def test_write_puts_name_and_picks_and_leaves_T_formula(sink):
    s, path = sink
    slot = s.next_slot()
    picks = list("LVLLLLLLE")
    s.write(slot, "CLEMENTE O. RRR", picks)
    ws = openpyxl.load_workbook(path)["ENTRY PG1"]
    r = slot.row
    assert ws.cell(r, 14).value == "CLEMENTE O. RRR"
    assert [ws.cell(r, 15 + k).value for k in range(9)] == picks
    assert str(ws.cell(r, 24).value).startswith("=COUNTIF")         # T untouched
    assert ws.cell(r, 13).value == 251                                # ticket # untouched
    assert s.next_slot().ticket == 252                                # advances


def test_blank_pick_leaves_cell_empty(sink):
    s, path = sink
    slot = s.next_slot()
    s.write(slot, "X", ["L", "", "E", "", "", "", "", "", "V"])
    ws = openpyxl.load_workbook(path)["ENTRY PG1"]
    assert [ws.cell(slot.row, 15 + k).value for k in range(9)] == ["L", None, "E", None, None, None, None, None, "V"]


def test_refuses_to_overwrite_then_allows(sink):
    s, _ = sink
    taken = s.slot_for_ticket(241)
    with pytest.raises(SlotTaken):
        s.write(taken, "NOPE", list("LLLLLLLLL"))
    s.write(taken, "FORCED", list("LLLLLLLLL"), allow_overwrite=True)


def test_undo_restores_empty(sink):
    s, path = sink
    slot = s.next_slot()
    rc = s.write(slot, "TEMP", list("EEEEEEEEE"))
    s.undo(rc)
    ws = openpyxl.load_workbook(path)["ENTRY PG1"]
    assert all(ws.cell(slot.row, 14 + k).value is None for k in range(10))
    assert s.next_slot().ticket == slot.ticket


def test_name_never_becomes_formula_or_date(sink):
    s, path = sink
    slot = s.next_slot()
    s.write(slot, "=1+1", list("LLLLLLLLL"))
    ws = openpyxl.load_workbook(path)["ENTRY PG1"]
    assert ws.cell(slot.row, 14).value == "=1+1" and ws.cell(slot.row, 14).data_type == "s"
    assert needs_text_prefix("1-2-3") and needs_text_prefix("=SUM(A1)") and not needs_text_prefix("1-2-3 RRR")


def test_format_name_uppercases_and_appends_suffix_once(sink):
    s, _ = sink
    assert s.format_name("clemente  o.", "RRR") == "CLEMENTE O. RRR"
    assert s.format_name("TONY 5000 RRR", "RRR") == "TONY 5000 RRR"


def test_full_block_is_reported(tmp_path):
    p = make_pool_workbook(tmp_path / "w.xlsx", filled=(52, 0, 0))
    s = ExcelSink(OpenpyxlBackend(str(p)))
    infos = s.block_infos()
    assert infos[0].free == 0 and infos[0].first_free is None
    assert s.default_block().block.name_col == 26                    # first block with room


def test_missing_header_is_a_clear_error(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "ENTRY PG1"
    wb.active["A1"] = "nothing here"
    wb.save(tmp_path / "bad.xlsx")
    with pytest.raises(ExcelError, match="header"):
        ExcelSink(OpenpyxlBackend(str(tmp_path / "bad.xlsx")))


def test_block_without_ticket_column():
    vals = [["NOMBRE"] + list(range(1, 10))] + [[None] * 10 for _ in range(5)]
    blocks = discover_blocks("S", Snapshot(1, 1, vals), rows_without_tickets=5)
    assert len(blocks) == 1 and blocks[0].ticket_col is None and blocks[0].rows == [2, 3, 4, 5, 6]


# ------------------------------------------------------------------ fake Excel COM object model
class _Rng:
    def __init__(self, ws, r1, c1, r2, c2):
        self.ws, self.r1, self.c1, self.r2, self.c2 = ws, r1, c1, r2, c2

    @property
    def Value(self):
        rows = tuple(tuple(self.ws.cells.get((r, c)) for c in range(self.c1, self.c2 + 1))
                     for r in range(self.r1, self.r2 + 1))
        return rows[0][0] if len(rows) == 1 and len(rows[0]) == 1 else rows

    @Value.setter
    def Value(self, v):
        for i, row in enumerate(v):
            for j, x in enumerate(row):
                if x is None:
                    self.ws.cells.pop((self.r1 + i, self.c1 + j), None)
                else:
                    self.ws.cells[(self.r1 + i, self.c1 + j)] = x


class _Cnt:
    def __init__(self, n):
        self.Count = n


class _Sheet:
    def __init__(self, name, cells):
        self.Name, self.cells = name, cells

    class _Cells:
        def __init__(s, ws): s.ws = ws
        def __call__(s, r, c): return (r, c)

    @property
    def Cells(self):
        return _Sheet._Cells(self)

    def Range(self, a, b):
        return _Rng(self, a[0], a[1], b[0], b[1])

    @property
    def UsedRange(self):
        rs = [k[0] for k in self.cells] or [1]
        cs = [k[1] for k in self.cells] or [1]
        class U:
            Row, Column = min(rs), min(cs)
            Rows, Columns = _Cnt(max(rs) - min(rs) + 1), _Cnt(max(cs) - min(cs) + 1)
        return U


class _Sheets:
    def __init__(self, sheets): self.sheets = sheets
    @property
    def Count(self): return len(self.sheets)
    def Item(self, i): return self.sheets[i - 1]
    def __call__(self, name): return next(s for s in self.sheets if s.Name == name)


class _Book:
    def __init__(self, name, sheets): self.Name, self.Worksheets = name, _Sheets(sheets)
    def Save(self): self.saved = True


class _Books:
    def __init__(self, books): self.books = books
    @property
    def Count(self): return len(self.books)
    def __call__(self, key):
        return self.books[key - 1] if isinstance(key, int) else next(b for b in self.books if b.Name == key)


class _App:
    def __init__(self, books):
        self.Workbooks = _Books(books)
        self.ActiveWorkbook = books[0]


def test_com_backend_against_fake_excel_object_model(tmp_path):
    src = openpyxl.load_workbook(make_pool_workbook(tmp_path / "src.xlsx"))
    cells = {}
    for ws_name in ("ENTRY PG1",):
        for row in src[ws_name].iter_rows():
            for c in row:
                if c.value is not None:
                    cells[(c.row, c.column)] = c.value if not str(c.value).startswith("=") else 0
    book = _Book("JORNADA 10 PC.xlsx", [_Sheet("ENTRY PG1", cells), _Sheet("RESULTS", {(1, 1): "r"})])
    other = _Book("OTHER.xlsx", [_Sheet("ENTRY PG1", {})])
    app = _App([book, other])
    be = ComBackend(app=app)
    assert be.workbook_name == "JORNADA 10 PC.xlsx" and "live" in be.label
    sink = ExcelSink(be)
    slot = sink.next_slot()
    assert slot.ticket == 251
    sink.write(slot, "CLEMENTE O. RRR", list("LVLLLLLLE"))
    assert cells[(slot.row, 14)] == "CLEMENTE O. RRR"
    assert [cells[(slot.row, 15 + k)] for k in range(9)] == list("LVLLLLLLE")
    assert sink.next_slot().ticket == 252
    assert ComBackend(workbook_hint="other", app=app).workbook_name == "OTHER.xlsx"
    with pytest.raises(ExcelError):
        ComBackend(workbook_hint="nope", app=app)
