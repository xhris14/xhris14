"""Start anywhere and continue from there.

The real sheet's first block (tickets 1-240) has no '#' in its header cell: it used to be read as a block without ticket
numbers, so the window opened on ticket 241. These tests pin down: the ticket column is found from the numbers under the
header; a block without any gets as many rows as its neighbours; entries can start at a typed ticket or at the cell clicked in
Excel and then CONTINUE from there; a row that already has data is skipped, not an error; and the end of the block is said."""
from __future__ import annotations

import types

import openpyxl
import pytest

import test_pipeline_gui as base
import test_wiring as w
from helpers import make_pool_workbook
from quiniela_scanner import gui
from quiniela_scanner.excel_sink import (BlockFull, ComBackend, ExcelSink, OpenpyxlBackend, SlotTaken, Snapshot,
                                         discover_blocks)

env, readings = base.env, base.readings
mk2 = w.mk2
inject, good, entered, named, settle = base.inject, base.good, base.entered, base.named, base.settle
ScriptedReader = base.ScriptedReader

USER = dict(first_tickets=(1, 241, 481), hash_headers=(False, True, True), filled=(0, 0, 0), rows=30)
T_COL, N_COL = 13, 14                           # first block of make_pool_workbook: '#' in M, NOMBRE in N


def user_sink(tmp_path, **kw):
    path = make_pool_workbook(tmp_path / "u.xlsx", **{**USER, **kw})
    return ExcelSink(OpenpyxlBackend(str(path))), path


# ======================================================================================= finding the ticket column
def test_a_first_block_without_a_hash_header_is_found_by_its_numbers(tmp_path):
    sink, _ = user_sink(tmp_path)
    infos = sink.block_infos()
    assert [i.block.ticket_col for i in infos] == [13, 25, 37]
    assert "tickets 1-30" in infos[0].label and infos[0].first_free.ticket == 1
    assert sink.default_block().block is infos[0].block and not sink.layout_warnings


def test_numbers_that_do_not_count_up_are_not_tickets():
    vals = [[None] * 3 for _ in range(6)]
    vals[1] = [None, "NOMBRE"] + list(range(1, 10))
    vals[2][0], vals[3][0] = 7, 7                                 # 7, 7: not 1, 2, 3
    (blk,) = discover_blocks("S", Snapshot(1, 1, [r + [None] * 9 for r in vals]), rows_without_tickets=4)
    assert blk.ticket_col is None and blk.rows == [3, 4, 5, 6]


def test_a_block_without_tickets_is_as_long_as_its_neighbours_and_says_so(tmp_path):
    path = make_pool_workbook(tmp_path / "t.xlsx", first_tickets=(1, 241, 481), filled=(0, 0, 0), rows=30,
                              hash_headers=(False, True, True))
    wb = openpyxl.load_workbook(path)
    ws = wb["ENTRY PG1"]
    for r in range(5, 40):
        ws.cell(r, 13).value = f"x{r}"                            # the first block's numbers are text: no ticket column
    wb.save(path)
    sink = ExcelSink(OpenpyxlBackend(str(path)))
    first = sink.block_infos()[-1]                                # blocks without tickets are listed last
    assert first.block.ticket_col is None and len(first.block.rows) == 30 and first.block.rows[0] == 5
    assert any("no ticket numbers found left of 'NOMBRE'" in x and "instead of by ticket" in x for x in sink.layout_warnings)
    assert sink.default_block().block.ticket_col is not None


# ======================================================================================= the sink
def test_slot_at_and_next_free_after(tmp_path):
    sink, _ = user_sink(tmp_path)
    s = sink.slot_at("ENTRY PG1", 4 + 5, N_COL + 4)               # a pick cell of ticket 5
    assert (s.ticket, s.row) == (5, 9) and sink.slot_at("ENTRY PG1", 9, T_COL).ticket == 5
    assert sink.slot_at("ENTRY PG1", 9, 11) is None and sink.slot_at("ENTRY PG1", 2, N_COL) is None
    assert sink.slot_at("ENTRY PG1", 9, N_COL + 10).ticket == 5   # its T column ...
    assert sink.slot_at("ENTRY PG1", 9, N_COL + 11).ticket == 245  # ... and the next column is the next block's '#'
    sink.write_row_fast(None, list("LEVLEVLEV"), "A", ticket=6)
    assert sink.next_free_after(s.block, 9).ticket == 7           # 6 is taken now
    assert sink.next_free_after(s.block, 34) is None


def test_or_next_skips_a_used_ticket_and_strict_refuses_it(tmp_path):
    sink, _ = user_sink(tmp_path)
    for t in (10, 11):
        sink.write_row_fast(None, list("LEVLEVLEV"), f"P{t}", ticket=t)
    with pytest.raises(SlotTaken):
        sink.write_row_fast(None, list("EEEEEEEEE"), "X", ticket=10)
    r = sink.write_row_fast(None, list("EEEEEEEEE"), "X", ticket=10, or_next=True)
    assert r.ticket == 12
    r = sink.write_row_fast(None, list("VVVVVVVVV"), "Y", ticket=12, or_next=True)
    assert r.ticket == 13


def test_or_next_at_the_end_of_the_block_is_block_full_with_a_clear_message(tmp_path):
    sink, _ = user_sink(tmp_path)
    for t in (29, 30):
        sink.write_row_fast(None, list("LEVLEVLEV"), f"P{t}", ticket=t)
    with pytest.raises(BlockFull, match="every row after ticket 29 already has data"):
        sink.write_row_fast(None, list("EEEEEEEEE"), "X", ticket=29, or_next=True)
    assert sink.write_row_fast(None, list("EEEEEEEEE"), "X").ticket == 1       # (plain 'next free' still fills the gaps)


def test_active_cell_through_com(tmp_path):
    import test_excel_fast as ef
    app, ws = ef.pool_app(tmp_path)
    be = ComBackend(app=app)
    app.ActiveSheet = ws
    app.ActiveCell = types.SimpleNamespace(Row=7, Column=14.0)
    assert be.active_cell() == ("ENTRY PG1", 7, 14)
    app.ActiveWorkbook = types.SimpleNamespace(Name="OTHER.xlsx")             # another workbook is in front
    assert be.active_cell() is None
    assert OpenpyxlBackend(str(make_pool_workbook(tmp_path / "f.xlsx"))).active_cell() is None


# ======================================================================================= the window
def ticket_row(backend, ticket, block=0):
    return w.row_of_ticket(backend, ticket, T_COL + 12 * block)


def test_the_window_opens_on_ticket_1(env, mk2):
    app, backend = mk2(reader=ScriptedReader(), wb=USER)
    assert "tickets 1-30" in app.block_var.get() and app.ticket_var.get() == "1"


def test_a_typed_ticket_is_where_the_stack_continues(env, mk2, readings):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader, wb=USER)
    app.ticket_var.set("10")
    assert app.ticket_custom and not app.ticket_follow
    for k in range(3):
        inject(app, reader, readings[k], good(f"A{k}"))
        assert entered(app, 1 + k)
    assert [j.ticket for j in app.entries] == [10, 11, 12]
    assert app.ticket_var.get() == "13" and app.ticket_custom and app.ticket_follow
    assert named(app, *app.entries)
    assert ticket_row(backend, 1)[1:] == [None] * 9                               # ticket 1 was never used
    assert ticket_row(backend, 12)[0] == "A2"
    app.undo_job(app.entries[2])                                                  # the last one is undone: its ticket is next again
    assert app.ticket_var.get() == "12" and app.ticket_custom
    inject(app, reader, readings[3], good("B"))
    assert entered(app, 4) and app.entries[3].ticket == 12


def test_a_row_somebody_filled_by_hand_is_skipped_while_continuing(env, mk2, readings):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader, wb=USER)
    app.ticket_var.set("10")
    inject(app, reader, readings[0], good("A"))
    assert entered(app, 1) and app.ticket_var.get() == "11"
    ws = backend._wbv["ENTRY PG1"]
    for wb in (backend._wb, backend._wbv):                                        # the operator types into tickets 11 and 12
        for r in range(1, 60):
            if wb["ENTRY PG1"].cell(r, T_COL).value in (11, 12):
                wb["ENTRY PG1"].cell(r, N_COL).value = "BY HAND"
    inject(app, reader, readings[1], good("B"))
    assert entered(app, 2) and not app.alert_var.get()
    assert app.entries[1].ticket == 13 and app.ticket_var.get() == "14"
    assert ws is backend._wbv["ENTRY PG1"]


def test_the_cell_clicked_in_excel_is_the_start(env, mk2, readings):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader, wb=USER)
    row = next(r for r in range(1, 60) if backend._wbv["ENTRY PG1"].cell(r, T_COL).value == 20)
    backend.active_cell = lambda: ("ENTRY PG1", row, N_COL + 3)
    app.start_at_cell()
    assert app.ticket_var.get() == "20" and app.ticket_custom and app.ticket_follow
    assert "Starting at ticket 20" in app.notice_var.get()
    inject(app, reader, readings[0], good("A"))
    assert entered(app, 1) and app.entries[0].ticket == 20 and app.ticket_var.get() == "21"
    assert "Starting at" not in app.notice_var.get()
    backend.active_cell = lambda: ("ENTRY PG1", row, N_COL + 3)                   # click on the row that is filled now
    app.start_at_cell()
    assert app.ticket_var.get() == "21" and "ticket 20 already has data" in app.notice_var.get()


def test_a_cell_outside_every_block_changes_nothing(env, mk2):
    app, backend = mk2(reader=ScriptedReader(), wb=USER)
    backend.active_cell = lambda: ("ENTRY PG1", 2, 3)
    app.start_at_cell()
    assert "not inside a block" in app.notice_var.get() and app.ticket_var.get() == "1" and not app.ticket_custom
    backend.active_cell = lambda: None
    app.start_at_cell()
    assert "No selected cell" in app.notice_var.get()
    backend.active_cell = lambda: (_ for _ in ()).throw(base.ExcelError("Excel error: boom"))
    app.start_at_cell()
    assert "Could not read the cell selected in Excel" in app.alert_var.get()


def test_the_end_of_the_block_is_announced_and_nothing_jumps_to_another_block(env, mk2, readings):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader, wb=USER)
    app.ticket_var.set("30")
    inject(app, reader, readings[0], good("A"))
    assert entered(app, 1) and app.entries[0].ticket == 30
    assert "end of the block you started in" in app.notice_var.get()
    assert not app.ticket_custom and app.ticket_var.get() == "1"                  # the block's own first free ticket
    inject(app, reader, readings[1], good("B"))
    assert entered(app, 2) and app.entries[1].ticket == 1 and "tickets 1-30" in app.block_var.get()


def test_a_ticket_typed_in_another_block_selects_that_block_and_continues_there(env, mk2, readings):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader, wb=USER)
    app.ticket_var.set("250")
    inject(app, reader, readings[0], good("A"))
    assert entered(app, 1) and app.entries[0].ticket == 250
    assert "tickets 241-270" in app.block_var.get() and app.ticket_var.get() == "251" and app.ticket_follow


def test_another_workbook_does_not_inherit_the_chosen_start(env, mk2, readings):
    reader = ScriptedReader()
    app, backend = mk2(reader=reader, wb=USER)
    app.ticket_var.set("10")
    inject(app, reader, readings[0], good("A"))
    assert entered(app, 1) and app.ticket_var.get() == "11" and app.ticket_follow
    app.connect_excel()                                                           # the same workbook: the start is kept
    assert app.ticket_var.get() == "11" and app.ticket_follow
    backend.workbook_name = "JORNADA 13 PC.xlsx"                                  # next jornada: start from its first free ticket
    app.connect_excel()
    assert app.ticket_var.get() == "1" and not app.ticket_custom and not app.ticket_follow


# ======================================================================================= an upright (portrait) monitor
@pytest.mark.parametrize("sw,sh,scale", [(1080, 1920, 1.0), (1080, 1920, 1.25)])
def test_an_upright_monitor_puts_the_card_under_the_picture_and_cuts_nothing_off(env, mk2, monkeypatch, sw, sh, scale):
    """The user's monitor stands upright (1080 px wide): the side-by-side layout needed 1214 px, so both edges were cut."""
    monkeypatch.setattr(gui.App, "_screen", lambda self: (sw, sh, scale))
    app, backend = mk2(reader=ScriptedReader(), wb=USER, wbname="JORNADA 3 CLAUSURA 2026.xlsx")
    settle(app, 0.4)
    W = app.winfo_width()
    assert app.narrow and app.stacked and W <= sw - 20 and app.winfo_reqwidth() <= W, (W, app.winfo_reqwidth())
    for name in ("tree", "undo_btn", "review_btn", "name_entry", "block_cb", "ticket_entry", "excel_lbl", "write_btn",
                 "preview", "stats_lbl"):
        w = getattr(app, name)
        assert w.winfo_viewable(), name
        assert w.winfo_rootx() - app.winfo_rootx() + w.winfo_width() <= W, name
    assert app.write_btn.winfo_rooty() > app.preview.winfo_rooty() + app.preview.winfo_height()     # the card is below
    assert not app._mid_sb.winfo_ismapped()                                         # and nothing needs scrolling


def test_a_wide_screen_keeps_the_side_by_side_layout(env, mk2, monkeypatch):
    monkeypatch.setattr(gui.App, "_screen", lambda self: (1920, 1080, 1.0))
    app, backend = mk2(reader=ScriptedReader(), wb=USER)
    settle(app, 0.2)
    assert not app.narrow and not app.stacked
    assert app.write_btn.winfo_rootx() >= app.preview.winfo_rootx() + app.preview.winfo_width()       # side by side
