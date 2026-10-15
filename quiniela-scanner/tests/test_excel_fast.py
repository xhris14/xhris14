"""The fast write path: one small read + one write per slip, a cache instead of a snapshot per slip.

The fake Excel object model below counts what a real Excel would be asked to do: *data operations* (Range.Value
get/set, Font.Size get/set, UsedRange, Save) and every COM *round trip* (also attribute lookups such as
Range(...) or .Font), and adds up a simulated latency from them.
"""
import random
import re
from collections import Counter

import openpyxl
import pytest

from quiniela_scanner.excel_sink import (N_PICKS, ComBackend, ExcelError, ExcelSink, OpenpyxlBackend, SlotTaken,
                                         Snapshot, col_letter)
from helpers import make_pool_workbook

NAME_COL = 14                       # block 1 of make_pool_workbook: '#' in M, NOMBRE in N, picks O..W
TICKET_COL = 13
LEV = "LEV"


# ------------------------------------------------------------------ instrumented fake Excel object model
class FakeComError(Exception):
    def __init__(self, hresult, msg="com_error"):
        super().__init__(msg)
        self.hresult = hresult


class Meter:
    def __init__(self, trip_ms=0.5, cell_ms=0.002):
        self.ops = Counter()                # data operations by kind
        self.trips = 0                      # every COM round trip
        self.sim_ms = 0.0                   # simulated latency
        self.trip_ms, self.cell_ms = trip_ms, cell_ms

    def trip(self, n=1, cells=0):
        self.trips += n
        self.sim_ms += n * self.trip_ms + cells * self.cell_ms

    def op(self, kind, cells=0, trips=1):
        self.ops[kind] += 1
        self.trip(trips, cells)

    def mark(self):
        return Counter(self.ops), self.trips, self.sim_ms

    def since(self, mark):
        ops0, trips0, ms0 = mark
        return Counter({k: v - ops0[k] for k, v in self.ops.items() if v - ops0[k]}), self.trips - trips0, \
            self.sim_ms - ms0


_ADDR = re.compile(r"^([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?$")


def _col_index(letters):
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def parse_areas(addr):
    areas = []
    for part in addr.split(","):
        m = _ADDR.match(part)
        assert m, f"bad address {part!r}"
        c1, r1 = _col_index(m[1]), int(m[2])
        c2, r2 = (_col_index(m[3]), int(m[4])) if m[3] else (c1, r1)
        areas.append((r1, c1, r2, c2))
    return areas


class FakeRangeFont:
    def __init__(self, rng):
        self.rng = rng

    @property
    def Size(self):
        cells = self.rng.cells()
        self.rng.ws.meter.op("font_get", len(cells))
        sizes = {self.rng.ws.fonts.get(c, 11) for c in cells}
        return sizes.pop() if len(sizes) == 1 else None            # Excel answers Null when the sizes differ

    @Size.setter
    def Size(self, v):
        cells = self.rng.cells()
        self.rng.ws.meter.op("font_set", len(cells))
        self.rng.ws.check_busy()
        for c in cells:
            self.rng.ws.fonts[c] = v


class FakeRange:
    def __init__(self, ws, areas):
        self.ws, self.areas = ws, areas

    def cells(self):
        return [(r, c) for r1, c1, r2, c2 in self.areas for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]

    @property
    def Font(self):
        self.ws.meter.trip()
        return FakeRangeFont(self)

    @property
    def Value(self):
        assert len(self.areas) == 1
        r1, c1, r2, c2 = self.areas[0]
        self.ws.meter.op("value_get", (r2 - r1 + 1) * (c2 - c1 + 1))
        self.ws.check_dead()
        rows = tuple(tuple(self.ws.cells.get((r, c)) for c in range(c1, c2 + 1)) for r in range(r1, r2 + 1))
        return rows[0][0] if len(rows) == 1 and len(rows[0]) == 1 else rows

    @Value.setter
    def Value(self, v):
        assert len(self.areas) == 1
        r1, c1, r2, c2 = self.areas[0]
        self.ws.meter.op("value_set", (r2 - r1 + 1) * (c2 - c1 + 1))
        self.ws.check_busy()
        grid = ((v,),) if not isinstance(v, tuple) else v
        assert len(grid) == r2 - r1 + 1 and all(len(row) == c2 - c1 + 1 for row in grid), "shape mismatch"
        for i, row in enumerate(grid):
            for j, x in enumerate(row):
                self.ws.put((r1 + i, c1 + j), x)


class FakeCellRef:
    def __init__(self, r, c):
        self.r, self.c = r, c


class FakeSheet:
    def __init__(self, name, cells, meter, formulas=()):
        self.Name, self.cells, self.meter = name, cells, meter
        self.fonts = {}                     # (r, c) -> size, default 11
        self.texts = set()                  # cells entered with the leading apostrophe (stored as text)
        self.formulas = set(formulas)       # cells Excel would have turned into a formula
        self.converted = set()              # strings Excel would have turned into a number / date
        self.dead = False                   # simulate a Worksheet object whose workbook was closed
        self.busy = 0                       # next N writes fail with "call rejected"

    # -- what Excel does with an assignment
    def put(self, key, x):
        self.texts.discard(key)
        self.formulas.discard(key)
        self.converted.discard(key)
        if x is None:
            self.cells.pop(key, None)
        elif isinstance(x, str) and x.startswith("'"):
            self.cells[key] = x[1:]
            self.texts.add(key)
        elif isinstance(x, str) and x.startswith("="):
            self.cells[key] = 0
            self.formulas.add(key)
        elif isinstance(x, str) and re.match(r"^[+\-@]|^[\d\s\-/.:,]+$", x):
            self.cells[key] = x
            self.converted.add(key)
        else:
            self.cells[key] = x

    def check_busy(self):
        self.check_dead()
        if self.busy > 0:
            self.busy -= 1
            raise FakeComError(-2147418111, "call was rejected by callee")

    def check_dead(self):
        if self.dead:
            raise FakeComError(-2147417848, "the object invoked has disconnected from its clients")

    # -- object model
    def Range(self, a, b=None):
        self.meter.trip()
        self.check_dead()
        if isinstance(a, str):
            return FakeRange(self, parse_areas(a))
        return FakeRange(self, [(a.r, a.c, b.r, b.c)])

    def Cells(self, r, c):
        self.meter.trip()
        return FakeCellRef(r, c)

    @property
    def UsedRange(self):
        self.meter.op("usedrange", trips=5)
        rs = [k[0] for k in self.cells] or [1]
        cs = [k[1] for k in self.cells] or [1]

        class Cnt:
            def __init__(s, n): s.Count = n

        class U:
            Row, Column = min(rs), min(cs)
            Rows, Columns = Cnt(max(rs) - min(rs) + 1), Cnt(max(cs) - min(cs) + 1)
        return U

    # -- what the USER does in Excel (not counted)
    def user_type(self, row, col, value):
        self.put((row, col), value)

    def _shift(self, from_row, delta):
        def moved(k):
            return (k[0] + delta, k[1]) if k[0] >= from_row else k
        self.cells = {moved(k): v for k, v in self.cells.items()}
        self.fonts = {moved(k): v for k, v in self.fonts.items()}
        self.texts = {moved(k) for k in self.texts}
        self.formulas = {moved(k) for k in self.formulas}

    def user_insert_row(self, row):
        self._shift(row, +1)

    def user_delete_row(self, row):
        self.cells = {k: v for k, v in self.cells.items() if k[0] != row}
        self._shift(row + 1, -1)

    def row(self, row, col=NAME_COL, n=N_PICKS + 1):
        return [self.cells.get((row, col + i)) for i in range(n)]


class FakeSheets:
    def __init__(self, meter, sheets):
        self.meter, self.sheets = meter, sheets

    @property
    def Count(self):
        return len(self.sheets)

    def Item(self, i):
        return self.sheets[i - 1]

    def __call__(self, name):
        self.meter.trip()
        return next(s for s in self.sheets if s.Name == name)


class FakeBook:
    def __init__(self, name, sheets, meter):
        self.Name, self.meter = name, meter
        self.Worksheets = FakeSheets(meter, sheets)

    def Save(self):
        self.meter.op("save")


class FakeBooks:
    def __init__(self, books, meter):
        self.books, self.meter = books, meter

    @property
    def Count(self):
        return len(self.books)

    def __call__(self, key):
        self.meter.trip()
        return self.books[key - 1] if isinstance(key, int) else next(b for b in self.books if b.Name == key)


class FakeApp:
    def __init__(self, books, meter):
        self.meter = meter
        self.Workbooks = FakeBooks(books, meter)
        self.ActiveWorkbook = books[0]
        self.screen_log = []                # every value ScreenUpdating was set to
        self.screen = True
        self.screen_during_writes = []      # ScreenUpdating at the moment Range.Value was assigned

    @property
    def ScreenUpdating(self):
        return self.screen

    @ScreenUpdating.setter
    def ScreenUpdating(self, v):
        self.meter.trip()
        self.screen = v
        self.screen_log.append(v)


def pool_app(tmp_path, meter=None, **pool_kw):
    """A fake Excel holding the pool workbook (tickets 241.. / 481.. / 721.., header row 4)."""
    meter = meter or Meter()
    src = openpyxl.load_workbook(make_pool_workbook(tmp_path / "src.xlsx", **pool_kw))["ENTRY PG1"]
    cells, formulas = {}, set()
    for row in src.iter_rows():
        for c in row:
            if c.value is None:
                continue
            if str(c.value).startswith("="):
                cells[(c.row, c.column)] = 0
                formulas.add((c.row, c.column))
            else:
                cells[(c.row, c.column)] = c.value
    pg1 = FakeSheet("ENTRY PG1", cells, meter, formulas)
    book = FakeBook("JORNADA 10 PC.xlsx", [pg1, FakeSheet("ENTRY PG2", {}, meter),
                                           FakeSheet("RESULTS", {(1, 1): "r"}, meter)], meter)
    return FakeApp([book], meter), pg1


def fresh(tmp_path, sink_kw=None, **pool_kw):
    app, ws = pool_app(tmp_path, **pool_kw)
    be = ComBackend(app=app)
    return ExcelSink(be, **(sink_kw or {})), be, app, ws


class Watch:
    """Differences of backend calls and of the fake's COM work between two points in time."""

    def __init__(self, be, app):
        self.be, self.app = be, app
        self.reset()

    def reset(self):
        self.calls0 = Counter(self.be.calls)
        self.mark0 = self.app.meter.mark()

    def calls(self):
        return Counter({k: v - self.calls0[k] for k, v in self.be.calls.items() if v - self.calls0[k]})

    def com(self):
        return self.app.meter.since(self.mark0)


def picks_of(i):
    return [LEV[(i + k) % 3] for k in range(9)]


# ------------------------------------------------------------------ the headline numbers
def test_sixty_slips_in_a_row_never_read_the_sheet_again(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    assert sink.refresh_count == 1 and be.calls["snapshot"] == 2          # ENTRY PG1 + ENTRY PG2, once
    start, snaps0 = app.meter.mark(), be.calls["snapshot"]
    watch = Watch(be, app)
    tickets, expected = [], {}
    for i in range(60):
        picks = picks_of(i)
        empties = i % 3 == 2
        if empties:
            picks[1] = picks[5] = picks[7] = ""                            # three boxes left empty -> F1 in size 9
        watch.reset()
        rc = sink.write_row_fast(None, picks, name=f"NAME {i}")
        calls, (ops, trips, _) = watch.calls(), watch.com()
        if empties:
            assert calls == Counter(read_row=1, font_sizes=1, write_row=1)
            assert ops == Counter(value_get=1, font_get=1, value_set=1, font_set=1) and trips == 12
        else:
            assert calls == Counter(read_row=1, write_row=1)
            assert ops == Counter(value_get=1, value_set=1) and trips == 6
        assert sum(calls.values()) <= 4 and sum(ops.values()) <= 4
        tickets.append(rc.ticket)
        expected[rc.ticket] = (f"NAME {i}", [p or "F1" for p in picks])
    ops, trips, ms = app.meter.since(start)
    assert ops["usedrange"] == 0 and be.calls["snapshot"] == snaps0 and sink.refresh_count == 1
    # block 1 had 42 free rows (tickets 251..292); the default block then moved on to block 2 (484..) by itself
    assert tickets == list(range(251, 293)) + list(range(484, 502))
    row_of = {v: k[0] for k, v in ws.cells.items() if k[1] in (TICKET_COL, TICKET_COL + 12)}
    for t, (name, picks) in expected.items():
        col = NAME_COL if t < 481 else NAME_COL + 12
        assert ws.row(row_of[t], col) == [name] + picks, t                 # each ticket got its own name + picks
    assert ws.cells[(row_of[483], NAME_COL + 12)] == "PLAYER 3 CH"         # block 2's earlier tickets untouched
    print(f"\n60 slips: {ops['value_get']} reads, {ops['value_set']} writes, {ops['font_get']} font reads, "
          f"{ops['font_set']} font writes, {ops['usedrange']} UsedRange reads, {trips} round trips, "
          f"{ms:.0f} simulated ms")


def test_per_slip_cost_with_pessimistic_latency_is_far_below_the_budget(tmp_path):
    meter = Meter(trip_ms=5.0, cell_ms=0.02)                # a sluggish 5 ms per COM round trip
    app, ws = pool_app(tmp_path, meter=meter)
    sink = ExcelSink(ComBackend(app=app))
    start = meter.mark()
    for i in range(20):
        picks = picks_of(i)
        picks[4] = ""
        sink.write_row_fast(None, picks, name=f"N{i}")
    ops, trips, ms = meter.since(start)
    assert ms / 20 < 150, f"{ms / 20:.1f} ms simulated per slip"
    assert ops["usedrange"] == 0 and ops["value_get"] == 20


def test_blocks_big_like_the_real_workbook_stay_cheap_in_python(tmp_path):
    import time
    meter = Meter()
    cells = {}
    for b in range(10):                                      # 10 blocks x 250 rows (120 columns of blocks)
        c0 = 1 + b * 12
        cells[(4, c0)], cells[(4, c0 + 1)] = "#", "NOMBRE"
        for k in range(9):
            cells[(4, c0 + 2 + k)] = k + 1
        for i in range(250):
            cells[(5 + i, c0)] = b * 1000 + i + 1
            if i < 120:
                cells[(5 + i, c0 + 1)] = f"P{i}"
                for k in range(9):
                    cells[(5 + i, c0 + 2 + k)] = "L"
    ws = FakeSheet("ENTRY BIG", cells, meter)
    app = FakeApp([FakeBook("BIG.xlsx", [ws], meter)], meter)
    sink = ExcelSink(ComBackend(app=app))
    t0 = time.perf_counter()
    for i in range(30):
        sink.write_row_fast(None, picks_of(i), name=f"X{i}")
        sink.block_infos()                                   # what the GUI does after every slip
    per_slip = (time.perf_counter() - t0) / 30
    assert per_slip < 0.1, f"{per_slip * 1000:.1f} ms of Python per slip"


# ------------------------------------------------------------------ ScreenUpdating / retry / dead sheet
def test_screen_updating_is_off_during_the_write_and_back_on_after(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    app.screen_log.clear()
    sink.write_row_fast(None, picks_of(0), name="A")
    assert app.screen_log == [False, True] and app.screen is True
    # even a failing write puts it back
    def boom():
        raise FakeComError(5, "boom")
    ws.check_busy = boom
    app.screen_log.clear()
    with pytest.raises(ExcelError, match="boom"):
        sink.write_row_fast(None, picks_of(1), name="B")
    assert app.screen_log and app.screen_log == [False, True] * (len(app.screen_log) // 2) and app.screen is True


def test_screen_updating_failures_are_ignored(tmp_path):
    sink, be, app, ws = fresh(tmp_path)

    class Broken(FakeApp):
        @property
        def ScreenUpdating(self):
            return True

        @ScreenUpdating.setter
        def ScreenUpdating(self, v):
            raise FakeComError(-2147418111, "busy")
    app.__class__ = Broken
    sink.write_row_fast(None, picks_of(0), name="A")                       # no exception
    assert ws.cells[(15, NAME_COL)] == "A"


def test_busy_excel_is_retried(tmp_path, monkeypatch):
    monkeypatch.setattr("quiniela_scanner.excel_sink.time.sleep", lambda s: None)
    sink, be, app, ws = fresh(tmp_path)
    ws.busy = 2                                                           # the first two writes bounce
    sink.write_row_fast(None, picks_of(0), name="A")
    assert ws.cells[(15, NAME_COL)] == "A" and app.screen is True
    ws.busy = 99
    with pytest.raises(ExcelError, match="busy"):
        sink.write_row_fast(None, picks_of(1), name="B")
    assert (16, NAME_COL) not in ws.cells


def test_a_dead_worksheet_object_is_resolved_again(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    sink.write_row_fast(None, picks_of(0), name="A")                       # sheet object now cached
    book = app.Workbooks(1)
    new = FakeSheet("ENTRY PG1", ws.cells, ws.meter, ws.formulas)           # workbook closed and reopened
    new.fonts, new.texts = ws.fonts, ws.texts
    book.Worksheets.sheets[0] = new
    ws.dead = True
    sink.write_row_fast(None, picks_of(1), name="B")
    assert new.cells[(16, NAME_COL)] == "B"


# ------------------------------------------------------------------ backends
def test_com_read_row_flattens_what_excel_returns(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    assert be.read_row("ENTRY PG1", 5, TICKET_COL, 1) == [241]                 # one cell: Excel gives a bare scalar
    assert be.read_row("ENTRY PG1", 5, TICKET_COL, 3) == [241, "PLAYER 1 RRR", "L"]   # ((a, b, c),) -> [a, b, c]
    assert be.read_row("ENTRY PG1", 15, NAME_COL, 2) == [None, None]
    assert be.read_row("ENTRY PG1", 15, TICKET_COL, 1) == [251]
    assert be.calls["read_row"] == 4


def test_backend_calls_counter_counts_each_public_method(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    base = Counter(be.calls)
    be.sheet_names()
    be.snapshot("ENTRY PG1")
    be.read_row("ENTRY PG1", 5, 13, 3)
    be.font_sizes("ENTRY PG1", 5, 15, 2)
    be.write_row("ENTRY PG1", 20, NAME_COL, ["Z"])
    be.save()
    for k in ("sheet_names", "snapshot", "read_row", "font_sizes", "write_row", "save"):
        assert be.calls[k] - base[k] == 1, k

    path = make_pool_workbook(tmp_path / "f.xlsx")
    fb = OpenpyxlBackend(str(path))
    fb.sheet_names()
    fb.snapshot("ENTRY PG1")
    assert fb.read_row("ENTRY PG1", 5, 13, 3) == [241, "PLAYER 1 RRR", "L"]
    fb.font_sizes("ENTRY PG1", 5, 15, 2)
    fb.write_row("ENTRY PG1", 20, NAME_COL, ["Z"])
    fb.save()
    assert fb.calls == Counter(sheet_names=1, snapshot=1, read_row=1, font_sizes=1, write_row=1, save=1)
    assert fb.read_row("ENTRY PG1", 20, NAME_COL, 2) == ["Z", None]            # sees the file as it is after the write


def test_com_write_groups_font_changes_into_one_multi_area_range(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    watch = Watch(be, app)
    be.write_row("ENTRY PG1", 20, NAME_COL, ["N", "L", "F1", "L", "F1", "F1", "L", "F1", "L", "L"], {2: 9, 4: 9, 5: 9, 7: 9})
    ops, trips, ms = watch.com()
    assert ops == Counter(value_set=1, font_set=1)                         # P20,R20:S20,U20 in a single assignment
    assert {(20, NAME_COL + o): ws.fonts[(20, NAME_COL + o)] for o in (2, 4, 5, 7)} == {
        (20, NAME_COL + o): 9 for o in (2, 4, 5, 7)}
    assert (20, NAME_COL + 1) not in ws.fonts and (20, NAME_COL + 3) not in ws.fonts
    be.write_row("ENTRY PG1", 20, NAME_COL, [None] * 10, {2: 11, 4: 12})
    assert ws.fonts[(20, NAME_COL + 2)] == 11 and ws.fonts[(20, NAME_COL + 4)] == 12


def test_com_font_read_is_one_range_read_when_the_sizes_are_uniform_and_per_cell_when_mixed(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    watch = Watch(be, app)
    assert be.font_sizes("ENTRY PG1", 20, NAME_COL + 1, 9) == [11] * 9
    assert watch.com()[0] == Counter(font_get=1)
    ws.fonts[(20, NAME_COL + 3)] = 10
    watch.reset()
    assert be.font_sizes("ENTRY PG1", 20, NAME_COL + 1, 9) == [11, 11, 10, 11, 11, 11, 11, 11, 11]
    assert watch.com()[0] == Counter(font_get=10)


# ------------------------------------------------------------------ what lands where
def test_name_and_picks_land_in_the_next_free_ticket_and_T_is_untouched(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    rc = sink.write_row_fast(sink.default_block().block, list("LVLLLLLLE"), name="CLEMENTE O. RRR")
    assert rc.ticket == 251 and rc.row == 15 and rc.col == rc.name_col == NAME_COL and rc.sheet == "ENTRY PG1"
    assert ws.row(15) == ["CLEMENTE O. RRR"] + list("LVLLLLLLE")
    assert rc.written == ["CLEMENTE O. RRR"] + list("LVLLLLLLE") and rc.previous == [None] * 10
    assert (15, NAME_COL + 10) in ws.formulas and ws.cells[(15, TICKET_COL)] == 251
    assert sink.next_slot().ticket == 252                                   # the cache moved on, no re-read
    assert sink.write_row_fast(sink.block_infos()[1], list("EEEEEEEEE"), "B").ticket == 484   # a BlockInfo works too


def test_empty_name_leaves_the_name_cell_empty(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    rc = sink.write_row_fast(None, list("LVLLLLLLE"))
    assert ws.cells.get((15, NAME_COL)) is None and rc.written[0] is None
    assert ws.row(15)[1:] == list("LVLLLLLLE")
    assert not sink.is_free(rc.slot)                                         # picks alone make the slot taken


def test_f1_marker_font_9_and_marker_off(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    rc = sink.write_row_fast(None, ["L", "", "E", "", "", "", "", "", "V"], name="X")
    assert ws.row(15)[1:] == ["L", "F1", "E", "F1", "F1", "F1", "F1", "F1", "V"]
    assert [ws.fonts.get((15, NAME_COL + 1 + k), 11) for k in range(9)] == [11, 9, 11, 9, 9, 9, 9, 9, 11]
    assert ws.fonts.get((15, NAME_COL), 11) == 11
    assert rc.prev_sizes == {2: 11, 4: 11, 5: 11, 6: 11, 7: 11, 8: 11}
    short = sink.write_row_fast(None, ["L", "E"], name="SHORT")              # missing trailing picks count as empty
    assert ws.row(short.row)[1:] == ["L", "E"] + ["F1"] * 7

    (tmp_path / "x").mkdir()
    off, be2, app2, ws2 = fresh(tmp_path / "x", sink_kw=dict(blank_marker=""))
    watch = Watch(be2, app2)
    off.write_row_fast(None, ["L", ""] + [""] * 7, name="PLAIN")
    assert ws2.row(15)[1:] == ["L"] + [None] * 8 and 15 not in {r for r, _ in ws2.fonts}
    assert watch.calls() == Counter(read_row=1, write_row=1)                 # no font work without a marker


@pytest.mark.parametrize("name", ["=1+1", "1-2-3", "12/10", "+5", "@HOME", "2024"])
def test_names_that_look_like_formulas_or_dates_are_stored_as_text(tmp_path, name):
    sink, be, app, ws = fresh(tmp_path)
    rc = sink.write_row_fast(None, picks_of(0), name=name)
    key = (rc.row, NAME_COL)
    assert ws.cells[key] == name and key in ws.texts
    assert key not in ws.formulas and key not in ws.converted
    rc2 = sink.write_row_fast(None, picks_of(1), name="")
    sink.update_name(rc2, name)                                              # the late name is protected as well
    key2 = (rc2.row, NAME_COL)
    assert ws.cells[key2] == name and key2 in ws.texts and key2 not in ws.formulas and key2 not in ws.converted


def test_write_wrapper_uses_the_fast_path(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    watch = Watch(be, app)
    slot = sink.next_slot()
    rc = sink.write(slot, "OLD API", list("LVLLLLLLE"))
    assert watch.calls() == Counter(read_row=1, write_row=1) and watch.com()[0]["usedrange"] == 0
    assert rc.ticket == 251 and ws.row(15) == ["OLD API"] + list("LVLLLLLLE")
    with pytest.raises(SlotTaken):
        sink.write(sink.slot_for_ticket(241), "NOPE", list("LLLLLLLLL"))
    assert ws.cells[(5, NAME_COL)] == "PLAYER 1 RRR"
    sink.write(sink.slot_for_ticket(241), "FORCED", list("LLLLLLLLL"), allow_overwrite=True)
    assert ws.cells[(5, NAME_COL)] == "FORCED"


# ------------------------------------------------------------------ the cache is stale
def test_a_name_typed_by_hand_in_the_next_free_row_is_not_overwritten(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    sink.write_row_fast(None, picks_of(0), name="A")                         # ticket 251
    ws.user_type(16, NAME_COL, "TYPED BY HAND")                              # ticket 252 is taken behind our back
    assert sink.next_slot().ticket == 252                                    # the cache does not know
    watch = Watch(be, app)
    rc = sink.write_row_fast(None, picks_of(1), name="B")
    assert rc.ticket == 253 and rc.row == 17
    assert ws.cells[(16, NAME_COL)] == "TYPED BY HAND" and all((16, NAME_COL + 1 + k) not in ws.cells for k in range(9))
    assert ws.row(17) == ["B"] + picks_of(1)
    assert watch.calls() == Counter(read_row=2, write_row=1)                 # one extra small read, nothing else
    assert sink.refresh_count == 1 and sink.next_slot().ticket == 254        # and the cache learned about row 16
    assert not sink.is_free(sink.slot_for_ticket(252))


def test_a_pick_typed_by_hand_in_a_free_row_also_counts_as_taken(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    ws.user_type(15, NAME_COL + 5, "V")
    rc = sink.write_row_fast(None, picks_of(0), name="A")
    assert rc.ticket == 252 and ws.cells[(15, NAME_COL + 5)] == "V" and (15, NAME_COL) not in ws.cells


def test_many_hand_typed_rows_are_skipped_then_one_refresh_finds_the_rest(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    for r in range(15, 45):                                                  # 30 rows typed behind our back
        ws.user_type(r, NAME_COL, f"HAND {r}")
    base_reads, base_snaps = be.calls["read_row"], be.calls["snapshot"]
    rc = sink.write_row_fast(None, picks_of(0), name="FOUND")
    assert rc.row == 45 and ws.cells[(45, NAME_COL)] == "FOUND"
    assert be.calls["read_row"] - base_reads == 25 + 1                       # 25 bounded probes, then one after refresh
    assert be.calls["snapshot"] - base_snaps == 2 and sink.refresh_count == 2   # exactly ONE full refresh
    assert all(ws.cells[(r, NAME_COL)] == f"HAND {r}" for r in range(15, 45))


@pytest.mark.parametrize("delta", [+1, -1])
def test_rows_inserted_or_deleted_above_move_the_ticket_numbers(tmp_path, delta):
    sink, be, app, ws = fresh(tmp_path)
    sink.write_row_fast(None, picks_of(0), name="A")                         # ticket 251 in row 15
    if delta > 0:
        ws.user_insert_row(2)                                                # everything moves down one row
    else:
        ws.user_delete_row(2)                                                # everything moves up one row
    before = dict(ws.cells)
    rc = sink.write_row_fast(None, picks_of(1), name="B")
    assert sink.refresh_count == 2
    assert rc.ticket == 252 and rc.row == 16 + delta
    assert ws.cells[(rc.row, TICKET_COL)] == 252                             # right ticket
    assert ws.row(rc.row) == ["B"] + picks_of(1)
    changed = {k for k in set(before) | set(ws.cells) if before.get(k) != ws.cells.get(k)}
    assert changed == {(rc.row, NAME_COL + k) for k in range(10)}            # nothing else was touched
    watch = Watch(be, app)
    assert sink.write_row_fast(None, picks_of(2), name="C").ticket == 253    # steady again: no more refreshes
    assert sink.refresh_count == 2 and watch.calls() == Counter(read_row=1, write_row=1)


def test_stale_ticket_with_an_explicit_ticket(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    ws.user_insert_row(1)
    rc = sink.write_row_fast(None, picks_of(0), name="E", ticket=260)
    assert rc.ticket == 260 and rc.row == 5 + 19 + 1 and ws.cells[(rc.row, TICKET_COL)] == 260
    assert sink.refresh_count == 2


def test_hand_cleared_row_is_reused_for_an_explicit_ticket_without_a_refresh(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    for k in range(10):
        ws.cells.pop((5, NAME_COL + k))
    assert not sink.is_free(sink.slot_for_ticket(241))                      # the cache still believes it is taken
    rc = sink.write_row_fast(None, picks_of(0), name="BACK", ticket=241)
    assert rc.row == 5 and sink.refresh_count == 1 and ws.cells[(5, NAME_COL)] == "BACK"


def test_explicit_ticket_taken_raises_slot_taken_and_overwrite_is_explicit(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    before = dict(ws.cells)
    with pytest.raises(SlotTaken, match="already has data"):
        sink.write_row_fast(None, list("LLLLLLLLL"), name="NOPE", ticket=241)
    assert ws.cells == before                                               # nothing written
    ws.user_type(15, NAME_COL + 2, "E")                                      # cache says 251 is free, Excel says no
    with pytest.raises(SlotTaken):
        sink.write_row_fast(None, list("LLLLLLLLL"), name="NOPE", ticket=251)
    assert not sink.is_free(sink.slot_for_ticket(251)) and sink.refresh_count == 1
    old = [ws.cells.get((5, NAME_COL + k)) for k in range(10)]
    rc = sink.write_row_fast(None, list("EEEEEEEEE"), name="FORCED", ticket=241, allow_overwrite=True)
    assert rc.previous == old and ws.row(5) == ["FORCED"] + list("EEEEEEEEE")
    sink.undo(rc)
    assert [ws.cells.get((5, NAME_COL + k)) for k in range(10)] == old


def test_unknown_ticket_is_a_clear_error_after_one_refresh(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    with pytest.raises(ExcelError, match="ticket 99999"):
        sink.write_row_fast(None, picks_of(0), name="X", ticket=99999)
    assert sink.refresh_count == 2


def test_full_block_raises_a_clear_error_after_one_refresh(tmp_path):
    sink, be, app, ws = fresh(tmp_path, filled=(52, 0, 0))
    full = sink.block_infos()[0].block
    assert sink.block_infos()[0].free == 0
    before = dict(ws.cells)
    with pytest.raises(ExcelError, match=r"no free ticket in block ENTRY PG1 column N \(tickets 241-292\)"):
        sink.write_row_fast(full, picks_of(0), name="X")
    assert ws.cells == before and sink.refresh_count == 2
    # the user frees a row meanwhile: the cache said "full", Excel has room -> found by that single refresh
    for k in range(10):
        ws.cells.pop((20, NAME_COL + k))
    rc = sink.write_row_fast(full, picks_of(0), name="ROOM")
    assert rc.row == 20 and rc.ticket == 256 and sink.refresh_count == 3
    # the default block (None) skips full blocks without any refresh
    assert sink.write_row_fast(None, picks_of(1), name="NEXT").ticket == 481


def test_all_blocks_full_is_a_clear_error(tmp_path):
    sink, be, app, ws = fresh(tmp_path, filled=(52, 52, 52))
    with pytest.raises(ExcelError, match="no free ticket in block"):
        sink.write_row_fast(None, picks_of(0), name="X")


def test_refresh_failure_keeps_the_old_cache(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    infos = sink.block_infos()
    ws.cells.pop((4, NAME_COL))                                              # the NOMBRE header of block 1 vanishes
    ws.cells.pop((4, NAME_COL + 12))
    ws.cells.pop((4, NAME_COL + 24))
    with pytest.raises(ExcelError, match="header"):
        sink.refresh()
    assert [i.label for i in sink.block_infos()] == [i.label for i in infos]


# ------------------------------------------------------------------ name later than the picks
def test_update_name_writes_only_the_name_cell(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    rc = sink.write_row_fast(None, list("LVLLLLLLE"), name="")
    watch = Watch(be, app)
    sink.update_name(rc, "CLEMENTE O. RRR")
    assert watch.calls() == Counter(read_row=1, write_row=1)
    ops, trips, ms = watch.com()
    assert ops == Counter(value_get=1, value_set=1) and ops["usedrange"] == 0
    assert ws.row(15) == ["CLEMENTE O. RRR"] + list("LVLLLLLLE")
    assert rc.written[0] == "CLEMENTE O. RRR"
    assert sink.refresh_count == 1


def test_update_name_replaces_a_placeholder_and_can_be_repeated(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    rc = sink.write_row_fast(None, picks_of(0), name="?")
    sink.update_name(rc, "JUAN")                                             # default expect = what we wrote ("?")
    sink.update_name(rc, "JUAN P.", expect="JUAN")
    sink.update_name(rc, "JUAN PEREZ")                                       # written[0] followed along
    assert ws.cells[(rc.row, NAME_COL)] == "JUAN PEREZ"
    assert ws.row(rc.row)[1:] == picks_of(0)


def test_update_name_guard_never_overwrites_what_the_user_typed(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    rc = sink.write_row_fast(None, picks_of(0), name="")
    ws.user_type(rc.row, NAME_COL, "MANUAL")                                 # user typed there meanwhile
    with pytest.raises(SlotTaken):
        sink.update_name(rc, "ROBOT")
    assert ws.cells[(rc.row, NAME_COL)] == "MANUAL" and rc.written[0] is None
    sink.update_name(rc, "ROBOT", expect="MANUAL")                           # unless the caller says it is that
    assert ws.cells[(rc.row, NAME_COL)] == "ROBOT"
    rc2 = sink.write_row_fast(None, picks_of(1), name="GUESS")
    ws.user_type(rc2.row, NAME_COL, "SOMETHING ELSE")
    with pytest.raises(SlotTaken):
        sink.update_name(rc2, "BETTER", expect="GUESS")
    assert ws.cells[(rc2.row, NAME_COL)] == "SOMETHING ELSE"


def test_names_applied_in_shuffled_order_land_in_the_right_rows(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    receipts = [sink.write_row_fast(None, picks_of(i), name="") for i in range(5)]
    assert [r.ticket for r in receipts] == [251, 252, 253, 254, 255]
    assert all(ws.cells.get((r.row, NAME_COL)) is None for r in receipts)
    order = list(range(5))
    random.Random(7).shuffle(order)
    assert order != sorted(order)
    watch = Watch(be, app)
    for i in order:
        sink.update_name(receipts[i], f"PERSON {i}")
    assert watch.calls() == Counter(read_row=5, write_row=5)
    for i, r in enumerate(receipts):
        assert ws.row(r.row) == [f"PERSON {i}"] + picks_of(i)
        assert ws.cells[(r.row, TICKET_COL)] == 251 + i
    assert sink.refresh_count == 1


def test_update_name_follows_its_ticket_when_rows_moved(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    rc = sink.write_row_fast(None, picks_of(0), name="")
    ws.user_insert_row(1)
    sink.update_name(rc, "LATE")
    assert rc.row == 16 and ws.cells[(16, NAME_COL)] == "LATE" and ws.cells[(16, TICKET_COL)] == 251
    assert (15, NAME_COL) in ws.cells and ws.cells[(15, NAME_COL)] != "LATE"
    for k, v in list(ws.cells.items()):                                      # the ticket is gone from the sheet
        if k[1] == TICKET_COL and v == 251:
            del ws.cells[k]
    with pytest.raises(ExcelError):
        sink.update_name(rc, "NOWHERE")


# ------------------------------------------------------------------ undo
def test_undo_restores_values_and_font_sizes_without_a_snapshot(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    ws.fonts[(15, NAME_COL + 2)] = 10                                       # one pick cell had its own size
    ws.fonts[(15, NAME_COL + 8)] = 12
    rc = sink.write_row_fast(None, ["L", "", "E", "", "", "V", "", "", "V"], name="TEMP")
    assert rc.prev_sizes == {2: 10, 4: 11, 5: 11, 7: 11, 8: 12}
    assert [ws.fonts.get((15, NAME_COL + 1 + k), 11) for k in range(9)] == [11, 9, 11, 9, 9, 11, 9, 9, 11]
    snaps = be.calls["snapshot"]
    watch = Watch(be, app)
    sink.undo(rc)
    assert watch.calls() == Counter(read_row=1, write_row=1) and be.calls["snapshot"] == snaps
    assert all(ws.cells.get((15, NAME_COL + k)) is None for k in range(10))
    assert [ws.fonts.get((15, NAME_COL + o), 11) for o in range(1, 10)] == [11, 10, 11, 11, 11, 11, 11, 12, 11]
    assert sink.next_slot().ticket == 251 and sink.is_free(rc.slot)           # the cache knows it is free again
    assert sink.refresh_count == 1


def test_undo_after_a_late_name_clears_everything(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    rc = sink.write_row_fast(None, picks_of(0), name="")
    sink.update_name(rc, "LATE")
    sink.undo(rc)
    assert all(ws.cells.get((15, NAME_COL + k)) is None for k in range(10))
    assert sink.write_row_fast(None, picks_of(1), name="AGAIN").ticket == 251


def test_undo_gives_back_what_an_overwrite_replaced(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    old = [ws.cells.get((6, NAME_COL + k)) for k in range(10)]
    rc = sink.write_row_fast(None, ["", "L"] + [""] * 7, name="X", ticket=242, allow_overwrite=True)
    assert ws.fonts[(6, NAME_COL + 1)] == 9 and (6, NAME_COL + 2) not in ws.fonts
    sink.undo(rc)
    assert [ws.cells.get((6, NAME_COL + k)) for k in range(10)] == old
    assert ws.fonts[(6, NAME_COL + 1)] == 11 and ws.cells[(6, TICKET_COL)] == 242


def test_undo_does_not_clobber_another_ticket_when_rows_moved(tmp_path):
    sink, be, app, ws = fresh(tmp_path)
    rc = sink.write_row_fast(None, picks_of(0), name="MINE")
    ws.user_insert_row(1)
    sink.undo(rc)
    assert ws.row(16) == [None] * 10 and ws.cells[(16, TICKET_COL)] == 251     # cleared in the moved row
    assert ws.cells[(15, TICKET_COL)] == 250 and ws.cells[(15, NAME_COL)] == "PLAYER 10 RRR"


# ------------------------------------------------------------------ save_after_write
def test_save_after_write_is_one_more_backend_call(tmp_path):
    app, ws = pool_app(tmp_path)
    be = ComBackend(app=app)
    sink = ExcelSink(be, save_after_write=True)
    watch = Watch(be, app)
    sink.write_row_fast(None, ["L", "", "E", "", "", "V", "", "", "V"], name="X")
    assert watch.calls() == Counter(read_row=1, font_sizes=1, write_row=1, save=1)
    assert sum(watch.calls().values()) == 4 and sum(watch.com()[0].values()) == 5


# ------------------------------------------------------------------ the file backend takes the same path
def test_openpyxl_backend_fast_path_roundtrip(tmp_path):
    path = make_pool_workbook(tmp_path / "JORNADA 10 PC.xlsx")
    be = OpenpyxlBackend(str(path))
    sink = ExcelSink(be)
    snaps = be.calls["snapshot"]
    rc = sink.write_row_fast(None, ["L", "", "V", "E", "L", "L", "L", "L", "L"], name="")
    sink.update_name(rc, "=1+1")
    ws = openpyxl.load_workbook(path)["ENTRY PG1"]
    assert ws.cell(15, NAME_COL).value == "=1+1" and ws.cell(15, NAME_COL).data_type == "s"
    assert [ws.cell(15, NAME_COL + 1 + k).value for k in range(9)] == ["L", "F1", "V", "E", "L", "L", "L", "L", "L"]
    assert ws.cell(15, NAME_COL + 2).font.sz == 9 and ws.cell(15, NAME_COL + 1).font.sz == 11
    assert str(ws.cell(15, NAME_COL + 10).value).startswith("=COUNTIF")
    assert be.calls["snapshot"] == snaps and be.calls["font_sizes"] == 1
    sink.undo(rc)
    ws = openpyxl.load_workbook(path)["ENTRY PG1"]
    assert all(ws.cell(15, NAME_COL + k).value is None for k in range(10)) and ws.cell(15, NAME_COL + 2).font.sz == 11


def test_openpyxl_backend_detects_a_row_typed_into_the_file(tmp_path):
    path = make_pool_workbook(tmp_path / "w.xlsx")
    be = OpenpyxlBackend(str(path))
    sink = ExcelSink(be)
    wb = openpyxl.load_workbook(path)
    wb["ENTRY PG1"].cell(15, NAME_COL).value = "FILE EDIT"
    wb.save(path)
    rc = sink.write_row_fast(None, picks_of(0), name="NEW")
    ws = openpyxl.load_workbook(path)["ENTRY PG1"]
    assert rc.ticket == 252 and ws.cell(15, NAME_COL).value == "FILE EDIT" and ws.cell(16, NAME_COL).value == "NEW"


def test_snapshot_set_grows_the_grid_and_ignores_cells_above_or_left():
    snap = Snapshot(3, 2, [[None, None], [None, None]])
    snap.set(4, 3, "x")
    snap.set(7, 6, "y")
    snap.set(1, 1, "outside")
    assert snap.get(4, 3) == "x" and snap.get(7, 6) == "y" and snap.last_row == 7 and snap.last_col == 6
    assert snap.get(1, 1) is None
    assert col_letter(NAME_COL) == "N"


def test_snapshot_blank_follows_the_same_rule_as_norm():
    snap = Snapshot(2, 3, [["", "  ", None, 0, None], ["L"], []])
    assert snap.blank(2, 3, 3)                                               # '', whitespace, None -> empty
    assert not snap.blank(2, 3, 4)                                           # a 0 is data
    assert not snap.blank(3, 3, 10) and snap.blank(3, 4, 10)                 # ragged row: cells past its end are empty
    assert snap.blank(4, 3, 10) and snap.blank(99, 3, 10) and snap.blank(1, 3, 10)   # rows outside the cache
    assert snap.blank(2, 1, 2)                                               # columns left of the cache


def test_a1_addresses_for_two_letter_columns_and_multi_area_ranges(tmp_path):
    from quiniela_scanner.excel_sink import _a1, _a1_areas
    assert _a1(15, 38) == "AL15" and _a1(15, 38, 10) == "AL15:AU15" and _a1(5, 1, 1) == "A5"
    assert _a1_areas(5, 14, [7, 2, 4, 5]) == "P5,R5:S5,U5" and _a1_areas(5, 14, [3]) == "Q5"
    sink, be, app, ws = fresh(tmp_path)
    third = sink.block_infos()[2].block                                      # NOMBRE in column AL
    assert col_letter(third.name_col) == "AL"
    rc = sink.write_row_fast(third, ["L", "", "V", "E", "L", "L", "L", "L", "L"], name="TRES")
    assert rc.ticket == 721 and ws.row(rc.row, third.name_col) == ["TRES", "L", "F1", "V", "E", "L", "L", "L", "L", "L"]
    assert ws.fonts[(rc.row, third.name_col + 2)] == 9
    sink.update_name(rc, "TRES OK")
    assert ws.cells[(rc.row, third.name_col)] == "TRES OK"
    sink.undo(rc)
    assert ws.row(rc.row, third.name_col) == [None] * 10 and ws.fonts[(rc.row, third.name_col + 2)] == 11
