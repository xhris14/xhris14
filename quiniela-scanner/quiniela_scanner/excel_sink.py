"""Getting scanned slips into the Excel pool sheet.

The sheet is laid out as side-by-side *ticket blocks*:

    #  | NOMBRE | 1 | 2 | ... | 9 | T        (header row)
    241| RECA 30| V | V | ... | L | 5        (one row per ticket)

``discover_blocks`` finds every block from its header (so nothing is hard-coded to a
column letter), ``ExcelSink`` picks the next free ticket and writes name + 9 picks.
T (and anything else to the right) is left alone - it is normally a formula.

Two backends:
  * ComBackend      - drives the Excel window you already have open (Windows, pywin32).
  * OpenpyxlBackend - edits an .xlsx file on disk (workbook must be closed in Excel).

Speed: the full-sheet snapshot is only a CACHE for choosing the next free ticket. It is taken once (and on
refresh()); a write then costs one small read (is the slot still free and still that ticket?) plus one write.
The name can be filled in later than the picks with update_name().

A row write is atomic from the caller's view: assigning the values is the commit point. Whatever comes after it
(font size, Save, ScreenUpdating) is best effort and reported in Receipt.warnings, never raised - an exception
from write_row_fast always means "nothing was entered".
"""
from __future__ import annotations

import fnmatch
import os
import re
import shutil
import tempfile
import time
from collections import Counter
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path

N_PICKS = 9


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def col_letter(n: int) -> str:
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def norm(v):
    """Cell value -> None | int | float | str (whole floats become ints)."""
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


# Excel error cells come back from COM as these plain negative ints (#NULL! .. #N/A, #GETTING_DATA).
EXCEL_ERRORS = range(-2146826288, -2146826244)
_ERROR_NAMES = {-2146826288: "#NULL!", -2146826281: "#DIV/0!", -2146826273: "#VALUE!", -2146826265: "#REF!",
                -2146826259: "#NAME?", -2146826252: "#NUM!", -2146826246: "#N/A", -2146826245: "#GETTING_DATA"}


def is_ticket(v) -> bool:
    """A usable ticket number: a real int - not a bool and not an Excel error code (#REF! is -2146826265)."""
    return isinstance(v, int) and not isinstance(v, bool) and v not in EXCEL_ERRORS


def _show(v) -> str:
    """A cell value for a message: errors by their Excel name, empty as 'empty'."""
    if v is None:
        return "empty"
    if isinstance(v, int) and v in EXCEL_ERRORS:
        return f"the Excel error {_ERROR_NAMES.get(v, v)}"
    return repr(v)


def _ticket_like(v) -> bool:
    """Looks like somebody meant it as a ticket number (an int, an error where one was expected, 260.5, '260 ')."""
    if isinstance(v, bool) or v is None:
        return False
    if isinstance(v, (int, float)):
        return True
    return isinstance(v, str) and bool(re.match(r"^\d+([.,]\d+)?$", v.strip()))


# What Range.Value = "..." does with a string: Excel parses it like typed input. Anything that would turn into a
# formula, number, date, time, percentage, currency, boolean or error value (or lose a leading apostrophe) is
# sent with a leading apostrophe. English and Spanish month names are covered because the user's Excel may be either.
_MON = (r"ENE(?:RO)?|FEB(?:RERO|RUARY)?|MAR(?:ZO|CH)?|ABR(?:IL)?|APR(?:IL)?|MAY(?:O)?|JUN(?:IO|E)?|JUL(?:IO|Y)?|"
        r"AGO(?:STO)?|AUG(?:UST)?|SEP(?:T(?:IEMBRE)?|TEMBER)?|SET(?:IEMBRE)?|OCT(?:UBRE|OBER)?|"
        r"NOV(?:IEMBRE|EMBER)?|DIC(?:IEMBRE)?|DEC(?:EMBER)?|JAN(?:UARY)?")
_PARSED_AS_VALUE = re.compile(
    r"^[=+\-@'#]"                                                          # formula / sign / ref / apostrophe / error
    r"|^[\d\s\-/.:,]+$"                                                    # 12/10  1-2-3  2024  5:30  1,5
    r"|^[\s$\u20ac\u00a3\u00a5(+\-]*\d[\d\s.,:/\-]*(?:E[+\-]?\d+)?[\s%)]*$"      # $5  5%  (5)  1E3  1.5.2020
    r"|^\d{1,2}(?::\d{2}){0,2}\s*[AP]\.?M\.?$"                              # 5 PM  10:30AM
    r"|^(?:TRUE|FALSE|VERDADERO|FALSO)$"
    rf"|^\d{{1,4}}[\s\-/.]*(?:{_MON})\.?(?:[\s\-/.,]*\d{{1,4}})?$"             # 5-MAR  5 MAR 2026
    rf"|^(?:{_MON})\.?[\s\-/.,]*\d{{1,4}}(?:[\s\-/.,]*\d{{1,4}})?$",          # MAR 5  MAR 2026  MAR 5, 2026
    re.I)


def needs_text_prefix(s: str) -> bool:
    """True if Excel would turn this string into a formula / number / date / boolean (so it needs the apostrophe)."""
    return bool(_PARSED_AS_VALUE.match(s))


def format_name(name: str, tag: str = "", position: str = "after", sep: str = " ",
                uppercase: bool = True) -> str:
    """Tidy the name and add the group tag before or after it.

    The tag is skipped only if the slip already carries it as a separate word at that end
    (so a tag of "E" is still added to "JOSE" or "ELENA")."""
    name = " ".join(name.split())
    tag = tag.strip()
    if uppercase:
        name, tag = name.upper(), tag.upper()
    if not tag or not name:
        return name
    words = name.split(" ")
    edge = words[0] if position == "before" else words[-1]
    if edge.upper() == tag.upper():
        return name
    return f"{tag}{sep}{name}" if position == "before" else f"{name}{sep}{tag}"


class ExcelError(Exception):
    pass


class SlotTaken(ExcelError):
    pass


class ExcelBusy(ExcelError):
    """Excel stayed busy (a cell in edit mode, a dialog, a long recalculation) for the whole retry budget.
    Nothing was written by the call that raised it; the caller may simply try again later."""


class BlockFull(ExcelError):
    """The selected block has no free ticket left (``.block`` is that Block, or None if it vanished)."""

    def __init__(self, message: str, block=None):
        super().__init__(message)
        self.block = block


# --------------------------------------------------------------------------- #
# Snapshot + layout discovery
# --------------------------------------------------------------------------- #
@dataclass
class Snapshot:
    top: int                       # 1-based row of values[0]
    left: int                      # 1-based column of values[0][0]
    values: list
    formula_cells: set = field(default_factory=set)    # (row, col) of formulas that SHOW nothing - never free
    notes: list = field(default_factory=list)          # what the backend had to leave out (columns / rows cap)

    def get(self, row: int, col: int):
        r, c = row - self.top, col - self.left
        if 0 <= r < len(self.values) and 0 <= c < len(self.values[r]):
            return norm(self.values[r][c])
        return None

    @property
    def last_row(self) -> int:
        return self.top + len(self.values) - 1

    @property
    def last_col(self) -> int:
        return self.left + (max(len(r) for r in self.values) if self.values else 0) - 1

    def blank(self, row: int, col: int, n: int) -> bool:
        """True if the n cells from (row, col) are all empty (same rule as norm(), without building values).
        A formula that currently shows '' is not empty: writing a value would destroy it."""
        if self.formula_cells and any((row, col + k) in self.formula_cells for k in range(n)):
            return False
        r, c = row - self.top, col - self.left
        if not 0 <= r < len(self.values):
            return True
        for v in self.values[r][max(c, 0):max(c + n, 0)]:
            if v is not None and not (isinstance(v, str) and not v.strip()):
                return False
        return True

    def set(self, row: int, col: int, value):
        """Update one cell of the cache in place (grows the grid if the cell is past its edge)."""
        r, c = row - self.top, col - self.left
        if r < 0 or c < 0:
            return
        while len(self.values) <= r:
            self.values.append([])
        line = self.values[r]
        if len(line) <= c:
            line.extend([None] * (c + 1 - len(line)))
        line[c] = value


@dataclass
class Block:
    sheet: str
    header_row: int
    name_col: int
    ticket_col: int | None
    rows: list
    first_ticket: int | None = None   # ticket of rows[0] when the block was found: survives rows inserted above it

    @property
    def pick_col(self) -> int:
        return self.name_col + 1

    @property
    def key(self) -> str:
        return f"{self.sheet}!{col_letter(self.name_col)}{self.header_row}"


@dataclass
class Slot:
    block: Block
    row: int
    ticket: int | None

    def describe(self) -> str:
        b = self.block
        t = f"ticket {self.ticket}" if self.ticket is not None else "next row"
        return f"{t}  ({b.sheet} row {self.row}, name in column {col_letter(b.name_col)})"


@dataclass
class BlockInfo:
    block: Block
    free: int
    total: int
    first_free: Slot | None
    suffix: str                    # dominant trailing word of existing names (e.g. "RRR")
    label: str


def _counts_up(snap: Snapshot, row: int, col: int) -> bool:
    """The cell at (row, col) and the one under it are ticket numbers one apart: 240, 241 / 1, 2."""
    a, b = snap.get(row, col), snap.get(row + 1, col)
    return is_ticket(a) and is_ticket(b) and b == a + 1


def discover_blocks(sheet: str, snap: Snapshot, name_header: str = "NOMBRE",
                    rows_without_tickets: int | None = None, max_header_row: int = 500,
                    skipped: list | None = None) -> list:
    """Every NOMBRE | 1..9 header of the snapshot as a Block. The ticket column below a header must hold real
    ticket numbers (is_ticket); the block ends at the first cell that does not. A header with no ticket number
    under it is not a block - it is described in ``skipped`` (if given) so the caller can say why.

    The ticket column is the one left of NOMBRE: headed '#' (or No / Ticket / Folio), or - for a block whose header
    cell is empty - holding two ticket numbers one apart right under the header (a header cell with any text, e.g. the
    'T' of the block before, is never taken: its scores can count up too). A block without any ticket column
    gets ``rows_without_tickets`` rows, by default as many as its ticketed neighbours on the sheet have (else 52)."""
    want = name_header.strip().upper()
    blocks = []
    last_hdr = min(snap.last_row, max_header_row)
    for r in range(snap.top, last_hdr + 1):
        line = snap.values[r - snap.top]
        for k, v in enumerate(line):
            if not isinstance(v, str) or v.strip().upper() != want:
                continue
            c = snap.left + k
            if any(str(snap.get(r, c + 1 + j)) != str(j + 1) for j in range(N_PICKS)):
                continue
            left = c - 1
            if str(snap.get(r, left) or "").upper() in ("#", "NO", "NO.", "TICKET", "FOLIO"):
                tc = left
            elif left >= snap.left and _blank_value(snap.get(r, left)) and _counts_up(snap, r + 1, left):
                tc = left                    # an EMPTY header cell (the first block of a sheet often has no '#'), 1, 2, 3 ... below it
            else:
                tc = None
            rows = []
            if tc is not None:
                rr = r + 1
                while is_ticket(snap.get(rr, tc)):
                    rows.append(rr)
                    rr += 1
            else:
                rows = [-1]                  # (filled in below: it depends on the other blocks of the sheet)
            if rows:
                blocks.append(Block(sheet, r, c, tc, rows, snap.get(rows[0], tc) if tc is not None else None))
            elif skipped is not None:
                skipped.append(f"{sheet} column {col_letter(c)}: the header '{name_header}' at row {r} has no ticket "
                               f"number under it (row {r + 1} is {_show(snap.get(r + 1, tc))}) - block not used")
    ticketed = [len(b.rows) for b in blocks if b.ticket_col is not None]
    for b in blocks:
        if b.ticket_col is None:
            n = rows_without_tickets or max(ticketed, default=52)
            b.rows = list(range(b.header_row + 1, b.header_row + 1 + n))
            if ticketed and skipped is not None:
                skipped.append(f"{sheet} column {col_letter(b.name_col)}: no ticket numbers found left of "
                               f"'{name_header}' (the cells there are {_show(snap.get(b.header_row, b.name_col - 1))} and "
                               f"{_show(snap.get(b.header_row + 1, b.name_col - 1))}) - entries go to the next free row "
                               f"of rows {b.rows[0]}-{b.rows[-1]} instead of by ticket")
    return blocks


def slot_is_free(snap: Snapshot, block: Block, row: int) -> bool:
    return snap.blank(row, block.name_col, N_PICKS + 1)


def _dominant_suffix(names: list) -> str:
    toks = [n.split()[-1] for n in names if isinstance(n, str) and len(n.split()) > 1]
    if len(toks) < 3:
        return ""
    word, cnt = Counter(toks).most_common(1)[0]
    return word if cnt / len(names) >= 0.5 and not word.isdigit() else ""


# --------------------------------------------------------------------------- #
# Backends
# --------------------------------------------------------------------------- #
def _a1(row: int, col: int, n: int = 1) -> str:
    """'N12' or, for n > 1 cells to the right, 'N12:W12'."""
    first = f"{col_letter(col)}{row}"
    return first if n == 1 else f"{first}:{col_letter(col + n - 1)}{row}"


def _a1_areas(row: int, col: int, offsets: list) -> str:
    """Multi-area address of the cells col+offset on one row: 'P5,R5:T5'."""
    runs, start, prev = [], None, None
    for off in sorted(offsets):
        if start is None:
            start = prev = off
        elif off == prev + 1:
            prev = off
        else:
            runs.append((start, prev))
            start = prev = off
    if start is not None:
        runs.append((start, prev))
    return ",".join(_a1(row, col + a, b - a + 1) for a, b in runs)


MAX_COLS = 256                     # hard cap on the width of a snapshot (a pool sheet is ~60 columns wide)
MAX_ROWS = 600


def _blank_value(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _grid(v) -> list:
    """What Range.Value / Range.Formula returns (scalar, or a tuple of row tuples) -> list of row lists."""
    if v is None:
        return [[None]]
    if not isinstance(v, tuple):
        return [[v]]
    return [list(r) for r in v]


class OpenpyxlBackend:
    """Edit an .xlsx on disk. Fine for testing / when Excel is closed; the workbook must NOT be
    open in Excel while writing. openpyxl may drop pictures/charts, so prefer ComBackend."""

    def __init__(self, path: str):
        import openpyxl
        self._openpyxl = openpyxl
        self.path = str(path)
        if Path(path).suffix.lower() in (".xlsm", ".xls", ".xlsb"):
            raise ExcelError(f"{Path(path).name}: only plain .xlsx files can be edited as a file (macros and "
                             f"other content would be lost) - open it in Excel and use the live connection")
        self.label = Path(path).name + " (file)"
        self.workbook_name = Path(path).name
        self.calls = Counter()                 # public-method call counts, for tests / diagnostics
        self._cached: dict = {}                # (sheet, row, col) -> the result Excel had cached for a formula
        self._load()
        self._remember_cached()

    def _load(self):
        self._load_formulas()
        self._load_values()

    def _load_formulas(self):
        self._wb = self._openpyxl.load_workbook(self.path)

    def _load_values(self):
        self._wbv = self._openpyxl.load_workbook(self.path, data_only=True)
        self._fill_cached()

    # openpyxl cannot write the results Excel cached next to a formula: after our first save every formula reads
    # as empty (a ticket column of =M5+1 would vanish). Keep the results from the file as Excel saved it and put
    # them back into the value view wherever the formula is still there.
    def _remember_cached(self):
        for ws in self._wb.worksheets:
            wv = self._wbv[ws.title]
            for row in ws.iter_rows(max_row=min(ws.max_row, MAX_ROWS), max_col=min(ws.max_column, MAX_COLS)):
                for cell in row:
                    if cell.data_type == "f":
                        v = wv.cell(cell.row, cell.column).value
                        if v is not None:
                            self._cached[(ws.title, cell.row, cell.column)] = v

    def _fill_cached(self):
        for (title, r, c), v in self._cached.items():
            if title in self._wb.sheetnames and self._wb[title].cell(r, c).data_type == "f":
                cell = self._wbv[title].cell(r, c)
                if cell.value is None:
                    cell.value = v

    def sheet_names(self):
        self.calls["sheet_names"] += 1
        self._load_formulas()
        return list(self._wb.sheetnames)

    def snapshot(self, sheet: str, max_rows: int = MAX_ROWS, max_cols: int = MAX_COLS) -> Snapshot:
        self.calls["snapshot"] += 1
        self._load_formulas()
        self._load_values()
        ws, wf = self._wbv[sheet], self._wb[sheet]
        nrows, ncols = min(ws.max_row, max_rows), min(ws.max_column, max_cols)
        notes = []
        if ws.max_column > max_cols:
            notes.append(f"{sheet}: the sheet is {ws.max_column} columns wide; only the first {max_cols} were read")
        if ws.max_row > max_rows:
            notes.append(f"{sheet}: the sheet has {ws.max_row} rows; only the first {max_rows} were read")
        rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=nrows, max_col=ncols, values_only=True)]
        formulas = set()                        # formulas that show nothing: writing there would destroy them
        for r, line in enumerate(wf.iter_rows(min_row=1, max_row=nrows, max_col=ncols), start=1):
            for c, cell in enumerate(line, start=1):
                if cell.data_type == "f" and _blank_value(rows[r - 1][c - 1]):
                    formulas.add((r, c))
        return Snapshot(1, 1, rows or [[None]], formulas, notes)

    def read_row(self, sheet: str, row: int, col: int, n: int) -> list:
        """n raw cell values of one row, starting at (row, col)."""
        self.calls["read_row"] += 1
        self._load_values()
        ws = self._wbv[sheet]
        return [ws.cell(row=row, column=col + i).value for i in range(n)]

    def active_cell(self):
        """A file has no selected cell."""
        return None

    def font_sizes(self, sheet: str, row: int, col: int, n: int) -> list:
        self.calls["font_sizes"] += 1
        self._load_formulas()
        ws = self._wb[sheet]
        return [ws.cell(row=row, column=col + i).font.sz for i in range(n)]

    def write_row(self, sheet: str, row: int, col: int, values: list, font_sizes: dict | None = None) -> list:
        """Write one row; returns warnings (always [] here: the file save is the only step and it raises)."""
        self.calls["write_row"] += 1
        self._load_formulas()
        ws = self._wb[sheet]
        for i, v in enumerate(values):
            cell = ws.cell(row=row, column=col + i)
            cell.value = v
            if isinstance(v, str):
                cell.data_type = "s"          # never let a name like "=1" become a formula
        for off, size in (font_sizes or {}).items():
            cell = ws.cell(row=row, column=col + off)
            f = copy(cell.font)
            f.sz = size
            cell.font = f
        self._save_file()
        return []

    def _save_file(self):
        """Write next to the workbook, then swap it in: a crash or a full disk mid-save leaves the old file."""
        tmp = ""
        try:
            fd, tmp = tempfile.mkstemp(prefix="~quiniela-", suffix=".tmp", dir=str(Path(self.path).resolve().parent))
            os.close(fd)
            self._wb.save(tmp)
            try:
                shutil.copymode(self.path, tmp)
            except OSError:
                pass
            os.replace(tmp, self.path)
        except OSError as e:
            raise ExcelError(f"could not save {self.workbook_name} (is it open in Excel?): {e}") from e
        finally:
            if tmp:
                try:
                    os.remove(tmp)
                except OSError:
                    pass

    def save(self):
        self.calls["save"] += 1


class ComBackend:
    """Drive the Excel instance that is already running (Windows + pywin32)."""

    BUSY = {-2147418111, -2147417846, -2146777998}   # call rejected / busy / VBA_E_IGNORE (0x800AC472: cell in edit mode)
    GONE = {-2147417848, -2147417851, -2147023174}   # disconnected / server fault / RPC server unavailable
    RETRY_TRIES = 6                                  # COM runs on the Tk thread: ~1.2 s at most, then the app
    RETRY_DELAY = 0.2                                # shows the problem instead of freezing
    SCREEN_TRIES = 2                                 # tries to put ScreenUpdating back (Excel may be in edit mode)

    def __init__(self, workbook_hint: str = "", app=None):
        if app is None:
            try:
                import pythoncom
                import win32com.client
            except ImportError as e:
                raise ExcelError("pywin32 is not installed (pip install pywin32)") from e
            try:
                pythoncom.CoInitialize()
            except Exception:
                pass
            try:
                app = win32com.client.GetActiveObject("Excel.Application")
            except Exception as e:
                raise ExcelError("Excel is not running - open your jornada workbook first "
                                 f"({type(e).__name__}: {e}). If it IS open, start Excel and this program with the "
                                 "same privileges (both normal, neither 'as administrator')") from e
        self.app = app
        self.calls = Counter()                 # public-method call counts, for tests / diagnostics
        self._sheets: dict = {}                # sheet name -> Worksheet COM object (saves 2 round trips per call)
        self._screen_stuck = False             # we turned ScreenUpdating off and could not turn it back on yet
        wb = self._retry(lambda: self._pick_workbook(workbook_hint))
        self.workbook_name = self._retry(lambda: str(wb.Name))
        self.label = f"{self.workbook_name} (Excel, live)"

    def _pick_workbook(self, hint: str):
        count = self.app.Workbooks.Count
        if count == 0:
            raise ExcelError("Excel is running but no workbook is open")
        if not hint:
            wb = self.app.ActiveWorkbook
            if wb is None:                      # pywin32's Nothing: only hidden workbooks (PERSONAL.XLSB) are loaded
                raise ExcelError("Excel has no active workbook window - click your jornada workbook "
                                 "(Protected View: press Enable Editing), then connect again")
            return wb
        for i in range(1, count + 1):
            wb = self.app.Workbooks(i)
            if hint.lower() in str(wb.Name).lower():
                return wb
        raise ExcelError(f"no open workbook matches '{hint}'")

    @property
    def wb(self):
        return self.app.Workbooks(self.workbook_name)

    def _ws(self, sheet: str):
        ws = self._sheets.get(sheet)
        if ws is None:
            ws = self._sheets[sheet] = self.wb.Worksheets(sheet)
        return ws

    @staticmethod
    def _codes(e) -> set:
        """Every code a com_error carries. An error raised INSIDE Excel arrives as DISP_E_EXCEPTION (-2147352567)
        with the callee's own code in excepinfo[5] - that is where 0x800AC472 (cell in edit mode) lives."""
        codes = {getattr(e, "hresult", None)}
        ei = getattr(e, "excepinfo", None)
        if isinstance(ei, (tuple, list)) and len(ei) >= 6:
            codes.add(ei[5])
        return codes

    @classmethod
    def _is_busy(cls, e) -> bool:
        """Excel (not us) refused for now."""
        return bool(cls._codes(e) & cls.BUSY) or (getattr(e, "hresult", None) is None and "busy" in str(e).lower())

    def _retry(self, fn):
        if self._screen_stuck:
            self._heal_screen()
        last = None
        dropped = False
        for _ in range(self.RETRY_TRIES):
            try:
                return fn()
            except ExcelError:                          # ours (no workbook, no active window ...): say it as is
                raise
            except Exception as e:                      # pywintypes.com_error
                if self._is_busy(e):
                    last = e
                    time.sleep(self.RETRY_DELAY)
                    continue
                if self._sheets and not dropped:        # a cached sheet may be dead (workbook reopened): resolve again
                    self._sheets.clear()
                    dropped = True
                    continue
                gone = bool(self._codes(e) & self.GONE)
                raise ExcelError(f"Excel error: {e}" + (" - Excel was closed or restarted? Press Connect / "
                                                        "reconnect Excel" if gone else "")) from e
        raise ExcelBusy("Excel is busy (is a cell being edited?). Press Esc in Excel and retry.") from last

    # ScreenUpdating: off while a row is written, then ALWAYS back to True (never "what it was": a False seen later
    # is a leftover - of a macro that died or of a failed restore - and a frozen-looking Excel is worse than a
    # repaint). If the restore is rejected (Excel went into edit mode meanwhile) it is retried by the next call.
    def _screen_off(self) -> bool:
        try:
            self.app.ScreenUpdating = False
        except Exception:
            return False
        self._screen_stuck = False
        return True

    def _screen_on(self) -> bool:
        for i in range(self.SCREEN_TRIES):
            try:
                self.app.ScreenUpdating = True
                self._screen_stuck = False
                return True
            except Exception:
                if i + 1 < self.SCREEN_TRIES:
                    time.sleep(self.RETRY_DELAY)
        self._screen_stuck = True
        return False

    def _heal_screen(self):
        try:
            self.app.ScreenUpdating = True
            self._screen_stuck = False
        except Exception:
            pass

    def sheet_names(self):
        self.calls["sheet_names"] += 1

        def go():
            ws = self.wb.Worksheets
            return [str(ws.Item(i).Name) for i in range(1, ws.Count + 1)]
        return self._retry(go)

    def snapshot(self, sheet: str, max_rows: int = MAX_ROWS, max_cols: int = MAX_COLS) -> Snapshot:
        """The cache the sink picks free tickets from. Width = the sheet's real last used column (Cells.Find over
        formulas, not the often inflated UsedRange), never more than max_cols; formulas that show nothing are
        noted from one Range.Formula read. Slow, so only at connect / refresh."""
        self.calls["snapshot"] += 1
        self._sheets.pop(sheet, None)                   # a (re)connect re-resolves the sheet object

        def go():
            ws = self._ws(sheet)
            ur = ws.UsedRange
            top, left = int(ur.Row), int(ur.Column)
            used_rows, used_cols = int(ur.Rows.Count), int(ur.Columns.Count)
            notes = []
            ncols = used_cols
            try:
                last = ws.Cells.Find(What="*", LookIn=-4123, LookAt=2, SearchOrder=2, SearchDirection=2)
                ncols = min(used_cols, 1 if last is None else max(int(last.Column) - left + 1, 1))
            except Exception as e:                      # Find is only an optimisation; the cap below still holds
                if self._is_busy(e):
                    raise
            if ncols > max_cols:
                notes.append(f"{sheet}: the sheet is {ncols} columns wide; only the first {max_cols} columns "
                             f"(to {col_letter(left + max_cols - 1)}) were read")
                ncols = max_cols
            nrows = min(used_rows, max_rows)
            if used_rows > max_rows:
                notes.append(f"{sheet}: the sheet has {used_rows} rows; only the first {max_rows} were read")
            rng = ws.Range(ws.Cells(top, left), ws.Cells(top + nrows - 1, left + ncols - 1))
            vals = _grid(rng.Value)
            formulas = set()
            try:
                for i, line in enumerate(_grid(rng.Formula)):
                    for j, f in enumerate(line):
                        if isinstance(f, str) and f[:1] == "=" and i < len(vals) and j < len(vals[i]) \
                                and _blank_value(vals[i][j]):
                            formulas.add((top + i, left + j))
            except Exception as e:
                if self._is_busy(e):
                    raise
                notes.append(f"{sheet}: could not read the formulas ({e}); formulas showing nothing look empty")
            return Snapshot(top, left, vals, formulas, notes)
        return self._retry(go)

    def read_row(self, sheet: str, row: int, col: int, n: int) -> list:
        """n raw cell values of one row, starting at (row, col), with ONE Range read."""
        self.calls["read_row"] += 1

        def go():
            v = self._ws(sheet).Range(_a1(row, col, n)).Value
            if isinstance(v, tuple):                    # n > 1 -> ((a, b, ...),)
                v = list(v[0]) if v and isinstance(v[0], tuple) else list(v)
            else:                                       # n == 1 -> the bare value
                v = [v]
            return (v + [None] * n)[:n]
        return self._retry(go)

    def active_cell(self):
        """(sheet, row, col) of the cell selected in THIS workbook's window; None when another workbook is in front
        or nothing usable is selected (a chart, say)."""
        self.calls["active_cell"] += 1

        def go():
            wb = self.app.ActiveWorkbook
            if wb is None or str(wb.Name) != self.workbook_name:
                return None
            cell, ws = self.app.ActiveCell, self.app.ActiveSheet
            if cell is None or ws is None:
                return None
            return str(ws.Name), int(cell.Row), int(cell.Column)
        return self._retry(go)

    def font_sizes(self, sheet: str, row: int, col: int, n: int) -> list:
        self.calls["font_sizes"] += 1

        def go():
            ws = self._ws(sheet)
            size = ws.Range(_a1(row, col, n)).Font.Size        # Excel answers Null when the sizes differ
            if size:
                return [size] * n
            return [ws.Range(_a1(row, col + i)).Font.Size for i in range(n)]
        return self._retry(go)

    def _text_safe(self, ws, row: int, col: int, values: list) -> list:
        """Strings Excel would parse (a name like 'MAR 5' or '=1+1') get a leading apostrophe - except in a cell
        formatted as Text, where the apostrophe would show. Only those rare cells cost an extra round trip."""
        out = list(values)
        for i, v in enumerate(values):
            if isinstance(v, str) and needs_text_prefix(v):
                try:
                    is_text = ws.Range(_a1(row, col + i)).NumberFormat == "@"
                except Exception as e:
                    if self._is_busy(e):
                        raise
                    is_text = False
                if not is_text:
                    out[i] = "'" + v
        return out

    def write_row(self, sheet: str, row: int, col: int, values: list, font_sizes: dict | None = None) -> list:
        """Write one row. Assigning the values is the commit point: if that raises, nothing was written. The font
        sizes are best effort - a failure there (protected sheet: 'Format cells' not allowed) is returned as a
        warning, never raised. -> list of warnings."""
        self.calls["write_row"] += 1
        warnings = []
        by_size: dict = {}
        for off, size in sorted((font_sizes or {}).items()):
            by_size.setdefault(size, []).append(off)

        def put():
            ws = self._ws(sheet)
            vals = self._text_safe(ws, row, col, values)
            rng = ws.Range(_a1(row, col, len(vals)))
            if len(vals) == 1:
                rng.Value = vals[0]
            else:
                rng.Value = (tuple(vals),)

        def resize(size, offs):
            self._ws(sheet).Range(_a1_areas(row, col, offs)).Font.Size = size

        screen_off = self._screen_off()
        try:
            self._retry(put)
            for size, offs in by_size.items():              # one multi-area range per distinct size
                try:
                    self._retry(lambda: resize(size, offs))
                except ExcelError as e:
                    warnings.append(f"font size {size} not applied in row {row} (protected sheet?): {e}")
        finally:
            if screen_off and not self._screen_on():
                warnings.append("Excel's ScreenUpdating could not be switched back on; it is retried by the next call")
        return warnings

    def save(self):
        self.calls["save"] += 1
        self._retry(lambda: self.wb.Save())


def open_backend(kind: str, workbook: str = "", path: str = ""):
    """kind: auto | com | file"""
    import sys
    if kind == "file" or (kind == "auto" and sys.platform != "win32"):
        if not path:
            raise ExcelError("backend 'file' needs workbook_path in config.json")
        return OpenpyxlBackend(path)
    return ComBackend(workbook)


# --------------------------------------------------------------------------- #
# The sink
# --------------------------------------------------------------------------- #
@dataclass
class Receipt:
    slot: Slot
    sheet: str
    row: int
    col: int                       # first column written = the NAME column
    previous: list                 # name + 9 picks as they were before (normalised), for undo
    written: list                  # name + 9 picks as written; written[0] is None while the name is still empty
    ticket: int | None
    prev_sizes: dict = field(default_factory=dict)     # font sizes to put back on undo: {offset: size}
    name_col: int = 0              # same as col; kept as its own field for callers that think in "name column"
    warnings: list = field(default_factory=list)       # best-effort steps that failed AFTER the values landed

    def __post_init__(self):
        if not self.name_col:
            self.name_col = self.col


MAX_PROBES = 25                    # free-looking slots tried (one small read each) before a full refresh
STOP_SCAN = 12                     # rows under a block looked at for ticket-like cells the block stopped before


def _block_stop(snap: Snapshot, b: Block, blocks=()):
    """None if the block simply ends. Else (row, value, k): its ticket column stops at `row` on a cell that is not
    a ticket number (text '260 ', 260.5, #REF!, a blank spacer) although k ticket-like cells follow within a
    few rows - the tickets below are NOT part of the block. The look-ahead ends at the header of a block that
    starts right underneath (its own tickets are not "cut off" ones)."""
    if b.ticket_col is None or not b.rows:
        return None
    stop = b.rows[-1] + 1
    end = min([x.header_row for x in blocks if x is not b and x.sheet == b.sheet and x.ticket_col == b.ticket_col
               and stop <= x.header_row < stop + STOP_SCAN] + [stop + STOP_SCAN])
    cells = [snap.get(r, b.ticket_col) for r in range(stop, end)]
    k = sum(1 for v in cells if _ticket_like(v))
    return (stop, cells[0], k) if k else None


def _stop_text(b: Block, stop: tuple) -> str:
    row, v, k = stop
    return (f"block ends at row {b.rows[-1]} because the ticket cell is not a number "
            f"(row {row}: {_show(v)}); {k} ticket-like cell(s) below it are NOT used")


def _formula_cells(snap: Snapshot, b: Block) -> list:
    """(row, col) of the name / pick cells of block b that hold a formula showing nothing (never overwritten)."""
    if not snap.formula_cells or not b.rows:
        return []
    return [(r, c) for (r, c) in snap.formula_cells
            if b.rows[0] <= r <= b.rows[-1] and b.name_col <= c <= b.name_col + N_PICKS]


def _survey(snaps: dict, blocks: list, skipped: list) -> tuple:
    """-> (warnings, duplicate tickets as {(sheet, ticket)}): what is wrong with the layout, in words."""
    out = [n for s in snaps.values() for n in s.notes] + list(skipped)
    seen: dict = {}
    for b in blocks:
        snap = snaps[b.sheet]
        stop = _block_stop(snap, b, blocks)
        if stop:
            out.append(f"{b.sheet} column {col_letter(b.name_col)}: {_stop_text(b, stop)}")
        for r in b.rows if b.ticket_col is not None else ():
            seen.setdefault((b.sheet, snap.get(r, b.ticket_col)), []).append(r)
        held = _formula_cells(snap, b)
        if held:
            out.append(f"{b.sheet} column {col_letter(b.name_col)}: {len(held)} name/pick cell(s) hold formulas that "
                       f"show nothing (first: row {min(r for r, _ in held)}) - those rows count as taken")
    dups = {k: rows for k, rows in seen.items() if len(rows) > 1}
    for (sheet, t), rows in list(dups.items())[:8]:
        out.append(f"{sheet}: ticket {t} appears {len(rows)} times (rows {', '.join(map(str, rows[:4]))}) - "
                   f"a name or Undo cannot tell those rows apart")
    if len(dups) > 8:
        out.append(f"... and {len(dups) - 8} more repeated ticket numbers")
    return out, set(dups)


class ExcelSink:
    def __init__(self, backend, sheet_patterns=("ENTRY*",), name_header="NOMBRE",
                 uppercase: bool = True, save_after_write: bool = False,
                 blank_marker: str = "F1", blank_font_size: int = 9):
        self.backend = backend
        self.blank_marker = blank_marker
        self.blank_font_size = blank_font_size
        self.patterns = list(sheet_patterns) or ["*"]
        self.name_header = name_header
        self.uppercase = uppercase
        self.save_after_write = save_after_write
        self.snaps: dict = {}              # CACHE of each ENTRY sheet; kept current in place by our own writes
        self.blocks: list = []
        self.layout_warnings: list = []    # what looks wrong in the sheet (filled by refresh(); show them once)
        self.dup_tickets: set = set()      # {(sheet, ticket)} appearing on more than one row
        self.refresh_count = 0             # full re-reads so far (construction counts as 1)
        self.refresh()

    # -- reading the layout ------------------------------------------------
    def refresh(self):
        """Re-read every ENTRY sheet in full. Slow over COM: only at start, on reconnect, or when a
        verification read finds the cache stale (write_row_fast does that by itself). Also rebuilds
        layout_warnings (truncated blocks, repeated / error tickets, formulas showing nothing, capped width)."""
        names = [n for n in self.backend.sheet_names()
                 if any(fnmatch.fnmatch(n.upper(), p.upper()) for p in self.patterns)]
        if not names:
            raise ExcelError(f"no sheet matches {self.patterns}; sheets are: {self.backend.sheet_names()}")
        snaps, blocks, skipped = {}, [], []
        for n in names:
            snap = self.backend.snapshot(n)
            snaps[n] = snap
            blocks += discover_blocks(n, snap, self.name_header, skipped=skipped)
        if not blocks:
            why = "; ".join(skipped + [x for s in snaps.values() for x in s.notes])
            raise ExcelError(f"no '{self.name_header} | 1 2 3 ... 9' header found on {names}" +
                             (f" ({why})" if why else ""))
        warnings, dups = _survey(snaps, blocks, skipped)
        self.snaps, self.blocks = snaps, blocks            # swapped only once everything was read
        self.layout_warnings, self.dup_tickets = warnings, dups
        self.refresh_count += 1
        return self.blocks

    def _ticket(self, block: Block, row: int):
        if block.ticket_col is None:
            return None
        t = self.snaps[block.sheet].get(row, block.ticket_col)
        return t if is_ticket(t) else None

    def block_infos(self) -> list:
        infos = []
        for b in self.blocks:
            snap = self.snaps[b.sheet]
            free = [r for r in b.rows if slot_is_free(snap, b, r)]
            tickets = [t for t in (self._ticket(b, r) for r in b.rows) if t is not None]
            names = [snap.get(r, b.name_col) for r in b.rows]
            names = [n for n in names if n]
            rng = f"tickets {min(tickets)}-{max(tickets)}" if tickets else f"rows {b.rows[0]}-{b.rows[-1]}"
            first = Slot(b, free[0], self._ticket(b, free[0])) if free else None
            suffix = _dominant_suffix(names)
            label = f"{b.sheet}  {col_letter(b.name_col)}:  {rng}  -  {len(free)} free"
            if suffix:
                label += f"  [{suffix}]"
            infos.append(BlockInfo(b, len(free), len(b.rows), first, suffix, label))
        infos.sort(key=lambda i: (i.block.ticket_col is None,
                                  self._ticket(i.block, i.block.rows[0]) or 0, i.block.name_col))
        return infos

    def default_block(self):
        infos = self.block_infos()
        return next((i for i in infos if i.free), infos[0] if infos else None)

    def next_slot(self, block: Block | None = None) -> Slot | None:
        for info in self.block_infos():
            if block is not None and info.block.key != block.key:
                continue
            if info.first_free:
                return info.first_free
        return None

    def slot_at(self, sheet: str, row: int, col: int) -> Slot | None:
        """The Slot (ticket row) of the block that covers this cell of the sheet - ticket, name, picks or T column -
        or None when the cell is outside every block. Free or not."""
        for b in self.blocks:
            if b.sheet != sheet or row not in b.rows:
                continue
            if (b.ticket_col if b.ticket_col is not None else b.name_col) <= col <= b.name_col + N_PICKS + 1:
                return Slot(b, row, self._ticket(b, row))
        return None

    def next_free_after(self, block: Block, row: int) -> Slot | None:
        """The first free Slot of the block below ``row`` (cache only), None when nothing is left below it."""
        b = self.resolve_block(block)
        if b is None:
            return None
        snap = self.snaps[b.sheet]
        r = next((r for r in b.rows if r > row and slot_is_free(snap, b, r)), None)
        return Slot(b, r, self._ticket(b, r)) if r is not None else None

    def active_cell(self):
        """(sheet, row, col) of the cell selected in the workbook's Excel window, None if the backend cannot say
        (file backend) or another workbook is in front."""
        fn = getattr(self.backend, "active_cell", None)
        return fn() if fn else None

    def slot_for_ticket(self, ticket: int) -> Slot | None:
        for b in self.blocks:
            for r in b.rows:
                if self._ticket(b, r) == ticket:
                    return Slot(b, r, ticket)
        return None

    def is_free(self, slot: Slot) -> bool:
        return slot_is_free(self.snaps[slot.block.sheet], slot.block, slot.row)

    # -- cache helpers (no COM) ---------------------------------------------
    def _first_free_row(self, block: Block, skip=(), after: int = 0):
        snap = self.snaps[block.sheet]
        return next((r for r in block.rows if r > after and r not in skip and slot_is_free(snap, block, r)), None)

    def resolve_block(self, old: Block | None) -> Block | None:
        """The Block of the CURRENT cache that stands for `old`, a Block taken from an earlier refresh (e.g. the one
        the operator selected). refresh() rebuilds the Blocks and rows inserted above a header change its key; the
        block is followed by sheet + column + first ticket, then by the nearest header. None for `old` = the default
        block; None is returned when `old` no longer exists (say so - do not fall back to another block)."""
        old = getattr(old, "block", old)                   # a BlockInfo works too
        if old is None:
            ordered = sorted(self.blocks, key=lambda b: (b.ticket_col is None,
                                                         self._ticket(b, b.rows[0]) or 0, b.name_col))
            return next((b for b in ordered if self._first_free_row(b) is not None),
                        ordered[0] if ordered else None)
        same =[b for b in self.blocks if b.sheet == old.sheet and b.name_col == old.name_col]
        if old.first_ticket is not None:
            same = [b for b in same if b.first_ticket == old.first_ticket] or same
        return next((b for b in same if b.key == old.key),
                    min(same, key=lambda b: abs(b.header_row - old.header_row), default=None))

    _resolve_block = resolve_block

    def _remember(self, sheet: str, row: int, col: int, values: list):
        snap = self.snaps[sheet]
        for k, v in enumerate(values):
            snap.set(row, col + k, v)

    # -- writing -----------------------------------------------------------
    def format_name(self, name: str, tag: str = "", position: str = "after", sep: str = " ") -> str:
        return format_name(name, tag, position, sep, self.uppercase)

    def _row_values(self, name: str, picks) -> tuple:
        """-> ([name, p1..p9], {offset: small font size}). A box left empty becomes the marker (default F1)."""
        fill = self.blank_marker or None
        letters = [p if p in ("L", "E", "V") else fill for p in picks]
        letters += [fill] * (N_PICKS - len(letters))
        values = [(name or "").strip() or None] + letters[:N_PICKS]
        small = {k: self.blank_font_size for k, v in enumerate(values) if k >= 1 and v is not None and v == fill}
        return values, small

    def _probe(self, slot: Slot) -> tuple:
        """ONE small read of [ticket?, name, 9 picks] -> (status, the 10 live cells), status being
        'stale' (the ticket number there is not the one we cached: rows were inserted/deleted),
        'taken' (some target cell has data, or holds a formula that shows nothing) or 'ok'.
        The cache learns what was read."""
        b = slot.block
        lead = 1 if b.ticket_col is not None else 0
        live = [norm(v) for v in self.backend.read_row(b.sheet, slot.row, b.name_col - lead, lead + N_PICKS + 1)]
        if lead and live[0] != slot.ticket:
            return "stale", live[lead:]
        cells = live[lead:]
        self._remember(b.sheet, slot.row, b.name_col, cells)
        fc = self.snaps[b.sheet].formula_cells           # known from the last full read: no extra round trip
        held = bool(fc) and any((slot.row, b.name_col + k) in fc for k in range(N_PICKS + 1))
        return ("taken" if held or any(v is not None for v in cells) else "ok"), cells

    def _landed(self, sheet: str, row: int, col: int, values: list) -> bool:
        """After a write that raised: does the row hold the intended values anyway? (one small read)"""
        want = [norm(v) for v in values]
        if all(v is None for v in want):
            return False
        return [norm(v) for v in self.backend.read_row(sheet, row, col, len(values))] == want

    def _write(self, sheet: str, row: int, col: int, values: list, small: dict | None) -> list:
        """backend.write_row -> its warnings. Raises only if the values did not go in; a write that raised although
        the row now holds the values counts as written (a retry would enter the slip twice)."""
        try:
            return list(self.backend.write_row(sheet, row, col, values, small or None) or [])
        except ExcelBusy:
            raise                                       # Excel rejected the call itself: nothing landed
        except ExcelError as e:
            try:
                landed = self._landed(sheet, row, col, values)
            except ExcelError:
                raise ExcelError(f"{e} - row {row} may already hold this data: look at it in Excel before entering "
                                 f"again") from e
            if not landed:
                raise
            return [f"Excel reported an error but the row holds the values: {e}"]

    def _save(self) -> list:
        """Save is best effort once the values are in the sheet (AutoSave normally covers it anyway)."""
        try:
            self.backend.save()
        except ExcelError as e:
            return [f"the values are in the sheet but Excel could not Save: {e}"]
        return []

    def _commit(self, slot: Slot, previous: list, values: list, small: dict) -> Receipt:
        b = slot.block
        prev_sizes = {}
        if small:                                       # old sizes only where the marker goes, in one read
            lo, hi = min(small), max(small)
            sizes = self.backend.font_sizes(b.sheet, slot.row, b.name_col + lo, hi - lo + 1)
            prev_sizes = {k: sizes[k - lo] for k in small if sizes[k - lo]}
        warnings = self._write(b.sheet, slot.row, b.name_col, values, small)
        self._remember(b.sheet, slot.row, b.name_col, values)      # the values are in: from here on nothing may raise
        fc = self.snaps[b.sheet].formula_cells
        gone = {(slot.row, b.name_col + k) for k in range(len(values))} & fc
        if gone:                                        # an explicit overwrite of a formula that showed nothing
            fc -= gone
            warnings.append(f"{len(gone)} cell(s) of row {slot.row} held a formula that is now replaced by the "
                            f"value (Undo cannot bring the formula back)")
        if self.save_after_write:
            warnings += self._save()
        return Receipt(slot, b.sheet, slot.row, b.name_col, list(previous), list(values), slot.ticket, prev_sizes,
                       warnings=warnings)

    def write_row_fast(self, block: Block | None, picks, name: str = "", ticket: int | None = None,
                       allow_overwrite: bool = False, or_next: bool = False) -> Receipt:
        """Write name + 9 picks into the next free ticket of `block` (None = default block), or into the
        explicit `ticket`. Typical cost: one small read + one write, no full-sheet read; a font read/write
        is added only for boxes left empty. `name` may be '' (the cell stays empty; fill it in later with
        update_name). SlotTaken only for an explicit ticket that already has data (unless allow_overwrite);
        a free-looking slot that turns out to be taken is skipped. allow_overwrite does not apply to "next free".
        or_next: an explicit ticket that has data is not an error - the next free ticket after it, in the same
        block, is used ("continue from there"; BlockFull if nothing is free below it).

        An exception means NOTHING was entered (ExcelBusy: Excel was busy; BlockFull: no free ticket left in the
        block; ExcelError: anything else). Once the values are in the sheet the call returns: font size / Save
        problems are in Receipt.warnings."""
        values, small = self._row_values(name, picks)
        block = getattr(block, "block", block)             # a BlockInfo works too
        if ticket is not None:
            return self._write_ticket(block, ticket, values, small, allow_overwrite, or_next)
        return self._write_next_free(block, values, small)

    def _write_next_free(self, block, values, small, after_ticket: int | None = None) -> Receipt:
        """First free row of the block - below the row of ``after_ticket`` when that is given."""
        stale = False
        for attempt in (0, 1):
            b = self.resolve_block(block)
            skip = set()
            stale = False
            after = 0
            if after_ticket is not None and b is not None:
                after = next((r for r in b.rows if self._ticket(b, r) == after_ticket), 0)
            for _ in range(MAX_PROBES):
                row = self._first_free_row(b, skip, after) if b else None
                if row is None:
                    break
                slot = Slot(b, row, self._ticket(b, row))
                status, live = self._probe(slot)
                if status == "ok":
                    return self._commit(slot, live, values, small)
                if status == "stale":
                    stale = True
                    break
                skip.add(row)                           # taken by hand; the cache now knows
            if attempt == 0:
                self.refresh()
        if stale:
            raise ExcelError("Excel changed while writing (ticket numbers moved twice); "
                             "nothing was written - try again")
        if b is None:
            raise ExcelError(f"no free ticket in block {block.key if block else '?'}: the block is gone from the sheet")
        first, last = b.rows[0], b.rows[-1]
        t = [self._ticket(b, first), self._ticket(b, last)]
        span = f"tickets {t[0]}-{t[1]}" if None not in t else f"rows {first}-{last}"
        snap = self.snaps[b.sheet]
        stop = _block_stop(snap, b, self.blocks)
        held = sorted({r for r, _ in _formula_cells(snap, b)})
        why = (f"every row after ticket {after_ticket} already has data" if after_ticket is not None
               else "every row already has data") + (f" or a formula that shows nothing ({len(held)} row(s), first row "
                                              f"{held[0]}; formulas are never overwritten)" if held else "")
        if stop:
            why += f"; {_stop_text(b, stop)}"
        raise BlockFull(f"no free ticket in block {b.sheet} column {col_letter(b.name_col)} ({span}): {why}", b)

    def _write_ticket(self, block, ticket, values, small, allow_overwrite, or_next=False) -> Receipt:
        for attempt in (0, 1):
            slot = self._find_ticket(ticket, block, only=or_next)     # (continuing: never in another block - another group)
            if slot is not None:
                status, live = self._probe(slot)
                if status == "ok" or (status == "taken" and allow_overwrite):
                    return self._commit(slot, live, values, small)
                if status == "taken" and or_next:
                    return self._write_next_free(slot.block, values, small, after_ticket=ticket)
                if status == "taken":
                    raise SlotTaken(f"{slot.describe()} already has data")
            if attempt == 0:
                self.refresh()
        if or_next and block is not None:
            raise ExcelError(f"ticket {ticket} is not in the selected block (any more) - pick the block or the ticket again")
        raise ExcelError(f"ticket {ticket} is not in the sheet (or its row keeps moving)")

    def _slots_for_ticket(self, ticket: int, block: Block | None = None, only: bool = False):
        """Slots holding ``ticket``: in ``block`` first, then in the other blocks - unless ``only`` (and a block is given)."""
        first = self.resolve_block(block) if block is not None else None
        others = [] if only and block is not None else [b for b in self.blocks if b is not first]
        for b in ([first] if first else []) + others:
            for r in b.rows:
                if self._ticket(b, r) == ticket:
                    yield Slot(b, r, ticket)

    def _find_ticket(self, ticket: int, block: Block | None = None, only: bool = False) -> Slot | None:
        return next(self._slots_for_ticket(ticket, block, only), None)

    def write(self, slot: Slot, name: str, picks: list, allow_overwrite: bool = False) -> Receipt:
        """Backwards-compatible: write into this slot (verified live, SlotTaken if it has data)."""
        values, small = self._row_values(name, picks)
        if slot.ticket is not None:
            return self._write_ticket(slot.block, slot.ticket, values, small, allow_overwrite)
        status, live = self._probe(slot)                 # a block without ticket numbers: nothing to relocate by
        if status == "taken" and not allow_overwrite:
            raise SlotTaken(f"{slot.describe()} already has data")
        return self._commit(slot, live, values, small)

    # -- changing a ticket after it was written -----------------------------
    def _picks_of(self, r: Receipt) -> list:
        return [norm(v) for v in r.written[1:]]

    def _read_receipt_row(self, r: Receipt, n: int) -> list:
        """The n cells from the name column of the receipt's row, after checking (same read) that the
        row still carries the receipt's ticket; if rows moved, refresh once and follow the ticket.
        A ticket number that appears on several rows cannot identify its row: then the picks must match too."""
        b = r.slot.block
        lead = 1 if b.ticket_col is not None and r.ticket is not None else 0
        dup = bool(lead) and (r.sheet, r.ticket) in self.dup_tickets
        span = max(n, N_PICKS + 1) if dup else n
        if lead + span == 0:
            return []
        for attempt in (0, 1):
            live = [norm(v) for v in self.backend.read_row(r.sheet, r.row, r.col - lead, lead + span)]
            if not lead or (live[0] == r.ticket and (not dup or live[lead + 1:] == self._picks_of(r))):
                return live[lead:lead + n]
            if attempt == 0:
                self.refresh()
                slot = None
                for cand in self._slots_for_ticket(r.ticket, b):
                    if not dup or [norm(v) for v in self.backend.read_row(
                            cand.block.sheet, cand.row, cand.block.name_col + 1, N_PICKS)] == self._picks_of(r):
                        slot = cand
                        break
                if slot is None:
                    break
                r.slot, r.sheet, r.row = slot, slot.block.sheet, slot.row
                r.col = r.name_col = slot.block.name_col
                b = slot.block
        raise ExcelError(f"ticket {r.ticket} moved or was removed in Excel; nothing was changed")

    def update_name(self, r: Receipt, name: str, expect=None) -> list:
        """Fill in / replace ONLY the name cell of a ticket written earlier (one write plus one small read).
        The cell must still be empty or hold `expect` (default: what we wrote there), else SlotTaken.
        -> warnings (also added to the receipt); raises only if the name did not go in."""
        want = norm(r.written[0] if expect is None else expect)
        cur = self._read_receipt_row(r, 1)[0]
        if cur is not None and cur != want:
            raise SlotTaken(f"{r.slot.describe()}: the name cell now holds {cur!r}; left as is")
        value = (name or "").strip() or None
        warnings = self._write(r.sheet, r.row, r.col, [value], None)
        self._remember(r.sheet, r.row, r.col, [value])
        r.written[0] = value
        if self.save_after_write:
            warnings += self._save()
        r.warnings += warnings
        return warnings

    def undo(self, r: Receipt) -> list:
        self._read_receipt_row(r, 0)                    # the row must still be that ticket (a read of its number only)
        warnings = self._write(r.sheet, r.row, r.col, list(r.previous), r.prev_sizes or None)
        self._remember(r.sheet, r.row, r.col, list(r.previous))
        if self.save_after_write:
            warnings += self._save()
        r.warnings += warnings
        return warnings
